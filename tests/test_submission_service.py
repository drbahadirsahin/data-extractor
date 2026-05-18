import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from dictionary_parser import FieldSpec
from run_extraction import ExtractionFieldResult, ExtractionResponse
from submission_service import (
    SubmissionValidationError,
    build_append_field_name_set,
    collect_unsubmittable_review_fields,
    build_redcap_record_payload,
    build_redcap_record_payload_rows,
    export_review_results_to_excel,
    normalize_tc_identity_no,
    repair_unsubmittable_fields_for_submission,
    resolve_patient_submission_plan,
    submit_patient_plan,
)
from workspace_extraction import PatientExtractionResult


class FakeIdentityClient:
    def __init__(self, lookup_result=None, create_result=None):
        self.lookup_result = lookup_result
        self.create_result = create_result

    def lookup_record_id(self, tc_identity_no: str):
        return self.lookup_result

    def create_record_by_tc(self, tc_identity_no: str):
        return self.create_result


class FakeRedcapClient:
    def __init__(self):
        self.calls = []

    def import_records(self, records, *, overwrite_behavior="normal", return_content="ids", force_auto_number=False):
        self.calls.append(
            {
                "records": records,
                "overwrite_behavior": overwrite_behavior,
                "return_content": return_content,
                "force_auto_number": force_auto_number,
            }
        )
        if force_auto_number:
            return ["501"]
        return []


class FakeProvider:
    def __init__(self, content: str) -> None:
        self.content = content
        self.calls = []

    def generate(self, messages, schema, settings):
        self.calls.append({"messages": messages, "schema": schema, "settings": settings})
        return type("LLMResponse", (), {"content": self.content})()


class SubmissionServiceTests(unittest.TestCase):
    def test_normalize_tc_identity_no(self):
        self.assertEqual(normalize_tc_identity_no("192 266 372 42"), "19226637242")
        self.assertIsNone(normalize_tc_identity_no("123"))

    def test_build_redcap_record_payload_prefers_code_for_radio(self):
        result = PatientExtractionResult(
            queue_label="hasta1",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="12",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="tani_ipss_skoru",
                        form_name="tan",
                        status="found",
                        final_value="Şiddetli",
                        final_value_code="3",
                    ),
                    ExtractionFieldResult(
                        field_name="hasta_ad",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="Ahmet",
                    ),
                ],
            ),
        )
        field_specs = {
            "tani_ipss_skoru": FieldSpec(
                field_name="tani_ipss_skoru",
                form_name="tan",
                field_type="radio",
                field_label="IPSS",
                choices="1, Hafif | 2, Orta | 3, Şiddetli",
            ),
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta adı",
            ),
        }
        payload, warnings = build_redcap_record_payload(
            result=result,
            target_record_id="12",
            field_specs_by_name=field_specs,
            excluded_field_names=set(),
        )
        self.assertEqual(payload["record_id"], "12")
        self.assertEqual(payload["tani_ipss_skoru"], "3")
        self.assertEqual(payload["hasta_ad"], "Ahmet")
        self.assertEqual(warnings, [])

    def test_build_redcap_record_payload_maps_choice_label_to_code(self):
        result = PatientExtractionResult(
            queue_label="hasta1",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="12",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="tani_ipss_skoru",
                        form_name="tan",
                        status="found",
                        final_value="Şiddetli",
                    ),
                ],
            ),
        )
        field_specs = {
            "tani_ipss_skoru": FieldSpec(
                field_name="tani_ipss_skoru",
                form_name="tan",
                field_type="radio",
                field_label="IPSS",
                choices="1, Hafif | 2, Orta | 3, Şiddetli",
            ),
        }

        payload, warnings = build_redcap_record_payload(
            result=result,
            target_record_id="12",
            field_specs_by_name=field_specs,
            excluded_field_names=set(),
        )

        self.assertEqual(payload["tani_ipss_skoru"], "3")
        self.assertEqual(warnings, [])

    def test_build_redcap_record_payload_skips_unmapped_choice_value(self):
        result = PatientExtractionResult(
            queue_label="hasta1",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="12",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="sonuc_durum",
                        form_name="sonuc",
                        status="uncertain",
                        final_value="remisyonda",
                        needs_review=True,
                        review_reasons=["choice_code_unmapped"],
                    ),
                    ExtractionFieldResult(
                        field_name="hasta_ad",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="Ahmet",
                    ),
                ],
            ),
        )
        field_specs = {
            "sonuc_durum": FieldSpec(
                field_name="sonuc_durum",
                form_name="sonuc",
                field_type="dropdown",
                field_label="Son durum",
                choices="1, Aktif | 2, Eksitus",
            ),
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta adı",
            ),
        }

        payload, warnings = build_redcap_record_payload(
            result=result,
            target_record_id="12",
            field_specs_by_name=field_specs,
            excluded_field_names=set(),
        )

        self.assertNotIn("sonuc_durum", payload)
        self.assertEqual(payload["hasta_ad"], "Ahmet")
        self.assertEqual(warnings, ["sonuc_durum: skipped_choice_code_unmapped"])

    def test_collect_unsubmittable_review_fields_tracks_multiple_patients(self):
        results = [
            PatientExtractionResult(
                queue_label="hasta1",
                patient_mode="existing",
                identifier_type="record_id",
                identifier_value="1",
                documents=[],
                document_results=[],
                merged_response=ExtractionResponse(
                    project_name="demo",
                    results=[
                        ExtractionFieldResult(
                            field_name="sonuc_durum",
                            form_name="sonuc",
                            status="uncertain",
                            final_value="remisyonda",
                            needs_review=True,
                            review_reasons=["choice_code_unmapped"],
                        ),
                    ],
                ),
            ),
            PatientExtractionResult(
                queue_label="hasta2",
                patient_mode="existing",
                identifier_type="record_id",
                identifier_value="2",
                documents=[],
                document_results=[],
                merged_response=ExtractionResponse(
                    project_name="demo",
                    results=[
                        ExtractionFieldResult(
                            field_name="sonuc_durum",
                            form_name="sonuc",
                            status="found",
                            final_value="Aktif",
                        ),
                    ],
                ),
            ),
        ]
        field_specs = {
            "sonuc_durum": FieldSpec(
                field_name="sonuc_durum",
                form_name="sonuc",
                field_type="dropdown",
                field_label="Son durum",
                choices="1, Aktif | 2, Eksitus",
            ),
        }

        issues = collect_unsubmittable_review_fields(results, field_specs)

        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].result_index, 0)
        self.assertEqual(issues[0].queue_label, "hasta1")
        self.assertEqual(issues[0].field_name, "sonuc_durum")

    def test_collect_unsubmittable_review_fields_tracks_invalid_dates(self):
        results = [
            PatientExtractionResult(
                queue_label="hasta1",
                patient_mode="existing",
                identifier_type="record_id",
                identifier_value="1",
                documents=[],
                document_results=[],
                merged_response=ExtractionResponse(
                    project_name="demo",
                    results=[
                        ExtractionFieldResult(
                            field_name="biyopsi_tarihi",
                            form_name="biyopsi",
                            status="found",
                            final_value=71,
                        ),
                    ],
                ),
            ),
        ]
        field_specs = {
            "biyopsi_tarihi": FieldSpec(
                field_name="biyopsi_tarihi",
                form_name="biyopsi",
                field_type="text",
                field_label="Biyopsi tarihi",
                text_validation="date_dmy",
            ),
        }

        issues = collect_unsubmittable_review_fields(results, field_specs)

        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].field_name, "biyopsi_tarihi")
        self.assertEqual(issues[0].reason, "invalid_date_format")

    def test_build_redcap_record_payload_normalizes_date_to_redcap_format(self):
        result = PatientExtractionResult(
            queue_label="hasta1",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="12",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="sonuc_son_izlem",
                        form_name="sonuc",
                        status="found",
                        final_value="16.01.2019",
                    ),
                ],
            ),
        )
        field_specs = {
            "sonuc_son_izlem": FieldSpec(
                field_name="sonuc_son_izlem",
                form_name="sonuc",
                field_type="text",
                field_label="Son izlem",
                text_validation="date_dmy",
            ),
        }

        payload, warnings = build_redcap_record_payload(
            result=result,
            target_record_id="12",
            field_specs_by_name=field_specs,
            excluded_field_names=set(),
        )

        self.assertEqual(payload["sonuc_son_izlem"], "2019-01-16")
        self.assertEqual(warnings, [])

    def test_build_redcap_record_payload_skips_invalid_date_values(self):
        result = PatientExtractionResult(
            queue_label="hasta1",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="12",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="hasta_kayit_tarihi",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value=1,
                    ),
                    ExtractionFieldResult(
                        field_name="biyopsi_tarihi",
                        form_name="biyopsi",
                        status="found",
                        final_value=71,
                    ),
                    ExtractionFieldResult(
                        field_name="hasta_ad",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="Ahmet",
                    ),
                ],
            ),
        )
        field_specs = {
            "hasta_kayit_tarihi": FieldSpec(
                field_name="hasta_kayit_tarihi",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Kayıt tarihi",
                text_validation="date_ymd",
            ),
            "biyopsi_tarihi": FieldSpec(
                field_name="biyopsi_tarihi",
                form_name="biyopsi",
                field_type="text",
                field_label="Biyopsi tarihi",
                text_validation="date_dmy",
            ),
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta adı",
            ),
        }

        payload, warnings = build_redcap_record_payload(
            result=result,
            target_record_id="12",
            field_specs_by_name=field_specs,
            excluded_field_names=set(),
        )

        self.assertNotIn("hasta_kayit_tarihi", payload)
        self.assertNotIn("biyopsi_tarihi", payload)
        self.assertEqual(payload["hasta_ad"], "Ahmet")
        self.assertEqual(
            warnings,
            [
                "hasta_kayit_tarihi: skipped_invalid_date_format",
                "biyopsi_tarihi: skipped_invalid_date_format",
            ],
        )

    def test_collect_unsubmittable_review_fields_tracks_invalid_numeric_values(self):
        results = [
            PatientExtractionResult(
                queue_label="hasta5",
                patient_mode="existing",
                identifier_type="record_id",
                identifier_value="5",
                documents=[],
                document_results=[],
                merged_response=ExtractionResponse(
                    project_name="demo",
                    results=[
                        ExtractionFieldResult(
                            field_name="rp_lnd_top_ln",
                            form_name="rp_patoloji",
                            status="found",
                            final_value="belirsiz",
                        ),
                    ],
                ),
            ),
        ]
        field_specs = {
            "rp_lnd_top_ln": FieldSpec(
                field_name="rp_lnd_top_ln",
                form_name="rp_patoloji",
                field_type="text",
                field_label="Toplam LN Sayısı",
                text_validation="integer",
            ),
        }

        issues = collect_unsubmittable_review_fields(results, field_specs)

        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].queue_label, "hasta5")
        self.assertEqual(issues[0].field_name, "rp_lnd_top_ln")
        self.assertEqual(issues[0].reason, "invalid_number_format")

    def test_build_redcap_record_payload_skips_invalid_numeric_values(self):
        result = PatientExtractionResult(
            queue_label="hasta5",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="5",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                            field_name="rp_lnd_top_ln",
                            form_name="rp_patoloji",
                            status="found",
                            final_value="belirsiz",
                    ),
                    ExtractionFieldResult(
                        field_name="hasta_ad",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="Ahmet",
                    ),
                ],
            ),
        )
        field_specs = {
            "rp_lnd_top_ln": FieldSpec(
                field_name="rp_lnd_top_ln",
                form_name="rp_patoloji",
                field_type="text",
                field_label="Toplam LN Sayısı",
                text_validation="integer",
            ),
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta adı",
            ),
        }

        payload, warnings = build_redcap_record_payload(
            result=result,
            target_record_id="5",
            field_specs_by_name=field_specs,
            excluded_field_names=set(),
        )

        self.assertNotIn("rp_lnd_top_ln", payload)
        self.assertEqual(payload["hasta_ad"], "Ahmet")
        self.assertEqual(warnings, ["rp_lnd_top_ln: skipped_invalid_number_format"])

    def test_build_redcap_record_payload_normalizes_absent_count_to_zero(self):
        result = PatientExtractionResult(
            queue_label="hasta5",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="5",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="rp_lnd_top_ln",
                        form_name="rp_patoloji",
                        status="found",
                        final_value="yok",
                    ),
                ],
            ),
        )
        field_specs = {
            "rp_lnd_top_ln": FieldSpec(
                field_name="rp_lnd_top_ln",
                form_name="rp_patoloji",
                field_type="text",
                field_label="Toplam LN Sayısı",
                text_validation="integer",
            ),
        }

        payload, warnings = build_redcap_record_payload(
            result=result,
            target_record_id="5",
            field_specs_by_name=field_specs,
            excluded_field_names=set(),
        )

        self.assertEqual(payload["rp_lnd_top_ln"], 0)
        self.assertEqual(warnings, [])

    def test_submission_value_repair_updates_absent_count_to_zero_before_submit(self):
        result = PatientExtractionResult(
            queue_label="hasta5",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="5",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="rp_lnd_top_ln",
                        form_name="rp_patoloji",
                        status="found",
                        final_value="yok",
                        needs_review=True,
                        review_reasons=["integer_validation_failed"],
                    ),
                ],
            ),
        )
        field_specs = {
            "rp_lnd_top_ln": FieldSpec(
                field_name="rp_lnd_top_ln",
                form_name="rp_patoloji",
                field_type="text",
                field_label="Toplam LN Sayısı",
                text_validation="integer",
            ),
        }

        repaired_count, warnings = repair_unsubmittable_fields_for_submission([result], field_specs)

        field_result = result.merged_response.results[0]
        self.assertEqual(repaired_count, 1)
        self.assertEqual(warnings, [])
        self.assertEqual(field_result.final_value, 0)
        self.assertFalse(field_result.needs_review)

    def test_submission_value_repair_uses_llm_for_remaining_invalid_numbers(self):
        result = PatientExtractionResult(
            queue_label="hasta5",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="5",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="rp_lnd_top_ln",
                        form_name="rp_patoloji",
                        status="found",
                        final_value="sıfır",
                    ),
                ],
            ),
        )
        field_specs = {
            "rp_lnd_top_ln": FieldSpec(
                field_name="rp_lnd_top_ln",
                form_name="rp_patoloji",
                field_type="text",
                field_label="Toplam LN Sayısı",
                text_validation="integer",
            ),
        }
        fake_provider = FakeProvider(
            '{"items":[{"item_id":"r1","normalized_value":0,"value_code":null,'
            '"confidence":0.92,"reason":"sıfır means zero"}]}'
        )

        with patch("submission_service.create_provider", return_value=fake_provider):
            repaired_count, warnings = repair_unsubmittable_fields_for_submission(
                [result],
                field_specs,
                llm_settings={"provider": "ollama", "model": "x"},
            )

        field_result = result.merged_response.results[0]
        self.assertEqual(repaired_count, 1)
        self.assertEqual(warnings, [])
        self.assertEqual(field_result.final_value, 0)
        self.assertFalse(field_result.needs_review)
        self.assertEqual(len(fake_provider.calls), 1)

    def test_build_redcap_record_payload_normalizes_valid_integer_text(self):
        result = PatientExtractionResult(
            queue_label="hasta5",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="5",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="rp_lnd_top_ln",
                        form_name="rp_patoloji",
                        status="found",
                        final_value="12.0",
                    ),
                ],
            ),
        )
        field_specs = {
            "rp_lnd_top_ln": FieldSpec(
                field_name="rp_lnd_top_ln",
                form_name="rp_patoloji",
                field_type="text",
                field_label="Toplam LN Sayısı",
                text_validation="integer",
            ),
        }

        payload, warnings = build_redcap_record_payload(
            result=result,
            target_record_id="5",
            field_specs_by_name=field_specs,
            excluded_field_names=set(),
        )

        self.assertEqual(payload["rp_lnd_top_ln"], 12)
        self.assertEqual(warnings, [])

    def test_resolve_new_patient_requires_duplicate_check(self):
        result = PatientExtractionResult(
            queue_label="yeni",
            patient_mode="new",
            identifier_type=None,
            identifier_value=None,
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="tc_kimlik_no",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="19226637242",
                    )
                ],
            ),
        )
        with self.assertRaises(SubmissionValidationError):
            resolve_patient_submission_plan(
                result=result,
                field_specs_by_name={},
                append_field_names=set(),
                repeating_forms=set(),
                repeating_events=set(),
                form_event_map={},
                identity_client=None,
            )

    def test_resolve_existing_patient_by_tc_uses_identity_lookup(self):
        result = PatientExtractionResult(
            queue_label="mevcut",
            patient_mode="existing",
            identifier_type="tc_kimlik_no",
            identifier_value="19226637242",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="hasta_ad",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="Ahmet",
                    )
                ],
            ),
        )
        plan = resolve_patient_submission_plan(
            result=result,
            field_specs_by_name={},
            append_field_names=set(),
            repeating_forms=set(),
            repeating_events=set(),
            form_event_map={},
            identity_client=FakeIdentityClient(lookup_result="77"),
        )
        self.assertEqual(plan.action, "update")
        self.assertEqual(plan.target_record_id, "77")
        self.assertEqual(plan.payload["record_id"], "77")

    def test_resolve_auto_patient_with_record_id_updates_without_identity_lookup(self):
        result = PatientExtractionResult(
            queue_label="auto-existing",
            patient_mode="auto",
            identifier_type="record_id",
            identifier_value="88",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="hasta_ad",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="Ahmet",
                    )
                ],
            ),
        )
        plan = resolve_patient_submission_plan(
            result=result,
            field_specs_by_name={},
            append_field_names=set(),
            repeating_forms=set(),
            repeating_events=set(),
            form_event_map={},
            identity_client=None,
        )
        self.assertEqual(plan.action, "update")
        self.assertEqual(plan.target_record_id, "88")

    def test_resolve_auto_patient_by_tc_updates_when_identity_exists(self):
        result = PatientExtractionResult(
            queue_label="auto-tc",
            patient_mode="auto",
            identifier_type="tc_kimlik_no",
            identifier_value="19226637242",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="hasta_ad",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="Ahmet",
                    )
                ],
            ),
        )
        plan = resolve_patient_submission_plan(
            result=result,
            field_specs_by_name={},
            append_field_names=set(),
            repeating_forms=set(),
            repeating_events=set(),
            form_event_map={},
            identity_client=FakeIdentityClient(lookup_result="77"),
        )
        self.assertEqual(plan.action, "update")
        self.assertEqual(plan.target_record_id, "77")

    def test_resolve_auto_patient_by_tc_creates_when_identity_missing(self):
        result = PatientExtractionResult(
            queue_label="auto-new",
            patient_mode="auto",
            identifier_type="tc_kimlik_no",
            identifier_value="19226637242",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="hasta_ad",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="Ahmet",
                    )
                ],
            ),
        )
        plan = resolve_patient_submission_plan(
            result=result,
            field_specs_by_name={},
            append_field_names=set(),
            repeating_forms=set(),
            repeating_events=set(),
            form_event_map={},
            identity_client=FakeIdentityClient(
                lookup_result=None,
                create_result=type("CreateResult", (), {"record_id": "9001", "created": True})(),
            ),
        )
        self.assertEqual(plan.action, "create")
        self.assertEqual(plan.target_record_id, "9001")

    def test_resolve_new_patient_uses_manual_tc_when_document_has_no_tc(self):
        result = PatientExtractionResult(
            queue_label="yeni",
            patient_mode="new",
            identifier_type=None,
            identifier_value=None,
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="hasta_ad",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="Ahmet",
                    )
                ],
            ),
        )
        plan = resolve_patient_submission_plan(
            result=result,
            field_specs_by_name={},
            append_field_names={"tc_kimlik_no"},
            repeating_forms=set(),
            repeating_events=set(),
            form_event_map={},
            identity_client=FakeIdentityClient(
                lookup_result=None,
                create_result=type("CreateResult", (), {"record_id": "9001", "created": True})(),
            ),
            manual_tc_identity_no="19226637242",
        )
        self.assertEqual(plan.action, "create")
        self.assertNotIn("tc_kimlik_no", plan.payload)
        self.assertEqual(plan.target_record_id, "9001")

    def test_export_review_results_to_excel_creates_workbook(self):
        result = PatientExtractionResult(
            queue_label="hasta1",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="12",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="hasta_soyad",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="AŞCI",
                    ),
                    ExtractionFieldResult(
                        field_name="hasta_ad",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="Ahmet",
                    )
                ],
            ),
            approved=True,
        )
        field_specs = {
            "hasta_soyad": FieldSpec(
                field_name="hasta_soyad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta soyadı",
            ),
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta adı",
            )
        }
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_path = Path(tmp_dir) / "review.xlsx"
            export_review_results_to_excel(output_path, [result], field_specs)
            workbook = load_workbook(output_path)
            self.assertEqual(workbook.sheetnames, ["FlatDataRaw", "FlatDataCode", "Patients", "Fields"])
            self.assertEqual(workbook["Patients"]["A2"].value, "hasta1")
            self.assertEqual(workbook["FlatDataRaw"]["A2"].value, "hasta1")
            headers = [cell.value for cell in workbook["FlatDataRaw"][1]]
            self.assertIn("hasta_ad", headers)
            self.assertLess(headers.index("hasta_soyad"), headers.index("hasta_ad"))

    def test_submit_patient_plan_create_uses_created_record_id_without_auto_number(self):
        plan = resolve_patient_submission_plan(
            result=PatientExtractionResult(
                queue_label="yeni",
                patient_mode="new",
                identifier_type=None,
                identifier_value=None,
                documents=[],
                document_results=[],
                merged_response=ExtractionResponse(
                    project_name="demo",
                    results=[
                        ExtractionFieldResult(
                            field_name="hasta_ad",
                            form_name="hasta_bilgileri",
                            status="found",
                            final_value="Ahmet",
                        )
                    ],
                ),
            ),
            field_specs_by_name={},
            append_field_names=set(),
            repeating_forms=set(),
            repeating_events=set(),
            form_event_map={},
            identity_client=FakeIdentityClient(
                lookup_result=None,
                create_result=type("CreateResult", (), {"record_id": "9001", "created": True})(),
            ),
            manual_tc_identity_no="19226637242",
        )
        identity_client = FakeIdentityClient(
            lookup_result=None,
            create_result=type("CreateResult", (), {"record_id": "9001", "created": True})(),
        )
        redcap_client = FakeRedcapClient()
        record_id = submit_patient_plan(
            plan=plan,
            redcap_client=redcap_client,
            identity_client=identity_client,
        )
        self.assertEqual(record_id, "9001")
        self.assertFalse(redcap_client.calls[0]["force_auto_number"])

    def test_build_append_field_name_set_collects_field_names(self):
        names = build_append_field_name_set(
            [{"Variable / Field Name": "tc_kimlik_no"}, {"Variable / Field Name": "ek_alan"}],
            {},
        )
        self.assertEqual(names, {"tc_kimlik_no", "ek_alan"})

    def test_build_redcap_record_payload_skips_append_fields(self):
        result = PatientExtractionResult(
            queue_label="hasta1",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="12",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="tc_kimlik_no",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="19226637242",
                    ),
                    ExtractionFieldResult(
                        field_name="hasta_ad",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="Ahmet",
                    ),
                ],
            ),
        )
        payload, _ = build_redcap_record_payload(
            result=result,
            target_record_id="12",
            field_specs_by_name={},
            excluded_field_names={"tc_kimlik_no"},
        )
        self.assertEqual(payload["record_id"], "12")
        self.assertEqual(payload["hasta_ad"], "Ahmet")
        self.assertNotIn("tc_kimlik_no", payload)

    def test_build_redcap_record_payload_rows_creates_repeating_row(self):
        result = PatientExtractionResult(
            queue_label="hasta1",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="12",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="lab_date",
                        form_name="tan_laboratuvar_sonucu",
                        status="found",
                        final_value="2025-10-03",
                    ),
                    ExtractionFieldResult(
                        field_name="lab_psa",
                        form_name="tan_laboratuvar_sonucu",
                        status="found",
                        final_value=4.04,
                    ),
                    ExtractionFieldResult(
                        field_name="hasta_ad",
                        form_name="hasta_bilgileri",
                        status="found",
                        final_value="Ahmet",
                    ),
                ],
            ),
        )
        field_specs = {
            "lab_date": FieldSpec(
                field_name="lab_date",
                form_name="tan_laboratuvar_sonucu",
                field_type="text",
                field_label="Lab tarihi",
            ),
            "lab_psa": FieldSpec(
                field_name="lab_psa",
                form_name="tan_laboratuvar_sonucu",
                field_type="text",
                field_label="PSA",
            ),
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta adı",
            ),
        }
        rows, warnings = build_redcap_record_payload_rows(
            result=result,
            target_record_id="12",
            field_specs_by_name=field_specs,
            excluded_field_names=set(),
            repeating_forms={"tan_laboratuvar_sonucu"},
            repeating_events=set(),
            form_event_map={"tan_laboratuvar_sonucu": ["event_1_arm_1"], "hasta_bilgileri": ["event_1_arm_1"]},
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["record_id"], "12")
        self.assertEqual(rows[0]["redcap_event_name"], "event_1_arm_1")
        self.assertEqual(rows[0]["hasta_ad"], "Ahmet")
        self.assertEqual(rows[1]["record_id"], "12")
        self.assertEqual(rows[1]["redcap_event_name"], "event_1_arm_1")
        self.assertEqual(rows[1]["redcap_repeat_instrument"], "tan_laboratuvar_sonucu")
        self.assertEqual(rows[1]["redcap_repeat_instance"], "new")
        self.assertEqual(rows[1]["lab_date"], "2025-10-03")
        self.assertEqual(rows[1]["lab_psa"], 4.04)
        self.assertEqual(warnings, [])

    def test_build_redcap_record_payload_rows_creates_repeating_event_row(self):
        result = PatientExtractionResult(
            queue_label="hasta1",
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value="12",
            documents=[],
            document_results=[],
            merged_response=ExtractionResponse(
                project_name="demo",
                results=[
                    ExtractionFieldResult(
                        field_name="biyopsi_tarihi",
                        form_name="klasik_biyopsi",
                        status="found",
                        final_value="2025-03-11",
                    ),
                ],
            ),
        )
        field_specs = {
            "biyopsi_tarihi": FieldSpec(
                field_name="biyopsi_tarihi",
                form_name="klasik_biyopsi",
                field_type="text",
                field_label="Biyopsi Tarihi",
            ),
        }
        rows, warnings = build_redcap_record_payload_rows(
            result=result,
            target_record_id="12",
            field_specs_by_name=field_specs,
            excluded_field_names=set(),
            repeating_forms=set(),
            repeating_events={"klasik_biyopsi_arm_1"},
            form_event_map={"klasik_biyopsi": ["klasik_biyopsi_arm_1"]},
        )
        self.assertEqual(
            rows,
            [
                {
                    "record_id": "12",
                    "redcap_event_name": "klasik_biyopsi_arm_1",
                    "redcap_repeat_instance": "new",
                    "biyopsi_tarihi": "2025-03-11",
                }
            ],
        )
        self.assertEqual(warnings, [])


if __name__ == "__main__":
    unittest.main()
