import tempfile
import unittest
import os
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from dictionary_parser import FieldSpec
from excel_import import ExcelColumnMapping, ExcelImportOptions, import_patient_excel
from llm_provider import LLMResponse
from settings_store import APP_HOME_ENV_VAR


class FakeProvider:
    def __init__(self, content: str) -> None:
        self.content = content
        self.calls = []

    def generate(self, messages, schema, settings):
        self.calls.append({"messages": messages, "schema": schema, "settings": settings})
        return LLMResponse(content=self.content)


class ExcelImportTests(unittest.TestCase):
    def test_import_patient_excel_maps_columns_and_choice_codes(self) -> None:
        field_specs = {
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta Adi",
            ),
            "tani_ipss_skoru": FieldSpec(
                field_name="tani_ipss_skoru",
                form_name="tani",
                field_type="radio",
                field_label="IPSS",
                choices="1, Hafif | 2, Orta | 3, Siddetli",
            ),
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "Hasta Adi", "IPSS", "Ignored"])
            sheet.append([12, "Ahmet", "Siddetli", "x"])
            sheet.append([13, "Ayse", "2", "y"])
            workbook.save(workbook_path)

            report = import_patient_excel(
                path=workbook_path,
                project_name="Demo",
                field_specs_by_name=field_specs,
                options=ExcelImportOptions(patient_id_column="record_id", patient_mode="existing"),
            )

        self.assertEqual(report.patient_id_column, "record_id")
        self.assertEqual([mapping.field_name for mapping in report.mapped_columns], ["hasta_ad", "tani_ipss_skoru"])
        self.assertEqual(report.ignored_columns, ["Ignored"])
        self.assertEqual(len(report.results), 2)
        first = report.results[0]
        self.assertEqual(first.identifier_type, "record_id")
        self.assertEqual(first.identifier_value, "12")
        fields = {item.field_name: item for item in first.merged_response.results}
        self.assertEqual(fields["hasta_ad"].final_value, "Ahmet")
        self.assertEqual(fields["tani_ipss_skoru"].final_value, "Siddetli")
        self.assertEqual(fields["tani_ipss_skoru"].final_value_code, "3")

        second_fields = {item.field_name: item for item in report.results[1].merged_response.results}
        self.assertEqual(second_fields["tani_ipss_skoru"].final_value, "Orta")
        self.assertEqual(second_fields["tani_ipss_skoru"].final_value_code, "2")

    def test_import_patient_excel_groups_rows_and_marks_conflicts(self) -> None:
        field_specs = {
            "tc_kimlik_no": FieldSpec(
                field_name="tc_kimlik_no",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="TC",
            ),
            "hasta_boy": FieldSpec(
                field_name="hasta_boy",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Boy",
                text_validation="integer",
                text_validation_min="75",
                text_validation_max="230",
            ),
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["tc_kimlik_no", "Boy"])
            sheet.append(["19226637242", 180])
            sheet.append(["19226637242", 181])
            workbook.save(workbook_path)

            report = import_patient_excel(
                path=workbook_path,
                project_name="Demo",
                field_specs_by_name=field_specs,
                options=ExcelImportOptions(patient_id_column="tc_kimlik_no", patient_mode="new"),
            )

        self.assertEqual(len(report.results), 1)
        patient = report.results[0]
        self.assertEqual(patient.patient_mode, "new")
        self.assertIsNone(patient.identifier_type)
        fields = {item.field_name: item for item in patient.merged_response.results}
        self.assertEqual(fields["tc_kimlik_no"].final_value, "19226637242")
        self.assertEqual(fields["hasta_boy"].status, "conflict")
        self.assertTrue(fields["hasta_boy"].needs_review)
        self.assertIn("multiple_documents_disagree", fields["hasta_boy"].review_reasons)

    def test_import_patient_excel_marks_unmapped_choice_for_review(self) -> None:
        field_specs = {
            "tani_ipss_skoru": FieldSpec(
                field_name="tani_ipss_skoru",
                form_name="tani",
                field_type="dropdown",
                field_label="IPSS",
                choices="1, Hafif | 2, Orta",
            ),
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "IPSS"])
            sheet.append(["77", "Bilinmeyen"])
            workbook.save(workbook_path)

            report = import_patient_excel(
                path=workbook_path,
                project_name="Demo",
                field_specs_by_name=field_specs,
            )

        result = report.results[0].merged_response.results[0]
        self.assertEqual(result.status, "uncertain")
        self.assertTrue(result.needs_review)
        self.assertIn("choice_code_unmapped", result.review_reasons)

    def test_import_patient_excel_can_normalize_invalid_choice_values_with_llm(self) -> None:
        field_specs = {
            "sonuc_durum": FieldSpec(
                field_name="sonuc_durum",
                form_name="sonuc",
                field_type="dropdown",
                field_label="Sonuç Durumu",
                choices="1, Remisyon | 2, Progresyon",
            ),
        }
        fake_provider = FakeProvider(
            '{"items":[{"item_id":"v1","normalized_value":"Remisyon","value_code":"1",'
            '"confidence":0.94,"reason":"remisyonda means remission"}]}'
        )

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "Sonuç Durumu"])
            sheet.append(["77", "remisyonda"])
            sheet.append(["78", "remisyonda"])
            workbook.save(workbook_path)

            with patch("excel_import.create_provider", return_value=fake_provider):
                report = import_patient_excel(
                    path=workbook_path,
                    project_name="Demo",
                    field_specs_by_name=field_specs,
                    options=ExcelImportOptions(
                        use_llm_value_normalization=True,
                        llm_settings={"provider": "ollama", "model": "x"},
                    ),
                )

        self.assertEqual(report.normalized_value_count, 2)
        self.assertEqual(len(fake_provider.calls), 1)
        self.assertEqual(fake_provider.calls[0]["messages"][1]["content"].count("remisyonda"), 1)
        for patient_result in report.results:
            fields = {item.field_name: item for item in patient_result.merged_response.results}
            self.assertEqual(fields["sonuc_durum"].final_value, "Remisyon")
            self.assertEqual(fields["sonuc_durum"].final_value_code, "1")
            self.assertFalse(fields["sonuc_durum"].needs_review)

    def test_import_patient_excel_normalizes_date_values_before_review(self) -> None:
        field_specs = {
            "biyopsi_tarihi": FieldSpec(
                field_name="biyopsi_tarihi",
                form_name="tani",
                field_type="text",
                field_label="Biyopsi Tarihi",
                text_validation="date_ymd",
            ),
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "Biyopsi Tarihi"])
            sheet.append(["77", "16.01.2019"])
            workbook.save(workbook_path)

            report = import_patient_excel(
                path=workbook_path,
                project_name="Demo",
                field_specs_by_name=field_specs,
            )

        result = report.results[0].merged_response.results[0]
        self.assertEqual(result.final_value, "2019-01-16")
        self.assertFalse(result.needs_review)

    def test_import_new_patient_adds_tc_identifier_from_selected_column(self) -> None:
        field_specs = {
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta Adi",
            ),
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Kimlik", "Hasta Adi"])
            sheet.append(["19226637242", "Ahmet"])
            workbook.save(workbook_path)

            report = import_patient_excel(
                path=workbook_path,
                project_name="Demo",
                field_specs_by_name=field_specs,
                options=ExcelImportOptions(
                    patient_id_column="Kimlik",
                    patient_mode="new",
                    identifier_type="tc_kimlik_no",
                ),
            )

        fields = {item.field_name: item for item in report.results[0].merged_response.results}
        self.assertEqual(fields["tc_kimlik_no"].final_value, "19226637242")
        self.assertEqual(fields["hasta_ad"].final_value, "Ahmet")

    def test_import_patient_excel_mixes_existing_record_id_and_auto_tc_patients(self) -> None:
        field_specs = {
            "tc_kimlik_no": FieldSpec(
                field_name="tc_kimlik_no",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="TC",
            ),
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta Adi",
            ),
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "tc_kimlik_no", "Hasta Adi"])
            sheet.append(["12", "", "Ahmet"])
            sheet.append(["", "19226637242", "Ayse"])
            workbook.save(workbook_path)

            report = import_patient_excel(
                path=workbook_path,
                project_name="Demo",
                field_specs_by_name=field_specs,
            )

        self.assertEqual(len(report.results), 2)
        self.assertEqual(report.results[0].patient_mode, "existing")
        self.assertEqual(report.results[0].identifier_type, "record_id")
        self.assertEqual(report.results[0].identifier_value, "12")
        self.assertEqual(report.results[1].patient_mode, "auto")
        self.assertEqual(report.results[1].identifier_type, "tc_kimlik_no")
        self.assertEqual(report.results[1].identifier_value, "19226637242")

    def test_import_patient_excel_uses_llm_column_mapping_for_custom_headers(self) -> None:
        field_specs = {
            "hasta_kilo": FieldSpec(
                field_name="hasta_kilo",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Kilo (kg)",
                text_validation="integer",
            ),
        }
        fake_provider = FakeProvider(
            '{"mappings":[{"column_name":"ClinicMetric42","field_name":"hasta_kilo","confidence":0.91,"reason":"same concept"}]}'
        )

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "ClinicMetric42"])
            sheet.append(["77", 82])
            workbook.save(workbook_path)

            with patch("excel_import.create_provider", return_value=fake_provider):
                report = import_patient_excel(
                    path=workbook_path,
                    project_name="Demo",
                    field_specs_by_name=field_specs,
                    options=ExcelImportOptions(use_llm_mapping=True, llm_settings={"provider": "ollama", "model": "x"}),
                )

        self.assertEqual([(item.column_name, item.field_name, item.match_type) for item in report.mapped_columns], [("ClinicMetric42", "hasta_kilo", "llm")])
        fields = {item.field_name: item for item in report.results[0].merged_response.results}
        self.assertEqual(fields["hasta_kilo"].final_value, 82)
        self.assertEqual(len(fake_provider.calls), 1)

    def test_import_patient_excel_allows_llm_split_full_name_column(self) -> None:
        field_specs = {
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Ad",
            ),
            "hasta_soyad": FieldSpec(
                field_name="hasta_soyad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Soyad",
            ),
        }
        fake_provider = FakeProvider(
            '{"mappings":['
            '{"column_name":"ClinicFullName","field_name":"hasta_ad","confidence":0.92,"value_transform":"name_given","reason":"full name first part"},'
            '{"column_name":"ClinicFullName","field_name":"hasta_soyad","confidence":0.91,"value_transform":"name_family","reason":"full name last part"}'
            ']}'
        )

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "ClinicFullName"])
            sheet.append(["77", "Ahmet Yılmaz"])
            workbook.save(workbook_path)

            with patch("excel_import.create_provider", return_value=fake_provider):
                report = import_patient_excel(
                    path=workbook_path,
                    project_name="Demo",
                    field_specs_by_name=field_specs,
                    options=ExcelImportOptions(use_llm_mapping=True, llm_settings={"provider": "ollama", "model": "x"}),
                )

        self.assertEqual(
            [(item.column_name, item.field_name, item.value_transform) for item in report.mapped_columns],
            [("ClinicFullName", "hasta_ad", "name_given"), ("ClinicFullName", "hasta_soyad", "name_family")],
        )
        fields = {item.field_name: item for item in report.results[0].merged_response.results}
        self.assertEqual(fields["hasta_ad"].final_value, "Ahmet")
        self.assertEqual(fields["hasta_soyad"].final_value, "Yılmaz")
        prompt_payload = fake_provider.calls[0]["messages"][1]["content"]
        self.assertIn("name_given", prompt_payload)
        self.assertIn("name_family", prompt_payload)

    def test_import_patient_excel_does_not_split_full_name_column_without_user_mapping(self) -> None:
        field_specs = {
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Ad",
            ),
            "hasta_soyad": FieldSpec(
                field_name="hasta_soyad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Soyad",
            ),
            "hasta_gobek_adi": FieldSpec(
                field_name="hasta_gobek_adi",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Göbek Adı",
            ),
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "İsim"])
            sheet.append(["77", "Ahmet Yılmaz"])
            workbook.save(workbook_path)

            report = import_patient_excel(
                path=workbook_path,
                project_name="Demo",
                field_specs_by_name=field_specs,
                options=ExcelImportOptions(use_llm_mapping=False),
            )

        self.assertEqual(report.mapped_columns, [])
        self.assertEqual(report.results[0].merged_response.results, [])

    def test_import_patient_excel_can_split_full_name_when_explicitly_allowed(self) -> None:
        field_specs = {
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Ad",
            ),
            "hasta_soyad": FieldSpec(
                field_name="hasta_soyad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Soyad",
            ),
        }

        from excel_import import build_procedural_column_mappings

        mappings = build_procedural_column_mappings(
            ["record_id", "İsim"],
            [{"record_id": "77", "İsim": "Ahmet Yılmaz"}],
            field_specs,
            excluded_columns={"record_id"},
            allow_composite_name_splits=True,
        )

        self.assertEqual(
            [(item.column_name, item.field_name, item.value_transform) for item in mappings],
            [("İsim", "hasta_ad", "name_given"), ("İsim", "hasta_soyad", "name_family")],
        )

    def test_import_patient_excel_manual_full_name_first_two_formulas(self) -> None:
        field_specs = {
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Ad",
            ),
            "hasta_soyad": FieldSpec(
                field_name="hasta_soyad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Soyad",
            ),
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "İsim"])
            sheet.append(["77", "Ahmet Yılmaz"])
            workbook.save(workbook_path)

            report = import_patient_excel(
                path=workbook_path,
                project_name="Demo",
                field_specs_by_name=field_specs,
                options=ExcelImportOptions(
                    use_llm_mapping=False,
                    automatic_mapping=False,
                    column_mappings=[
                        ExcelColumnMapping(
                            column_name="İsim",
                            field_name="hasta_ad",
                            match_type="manual",
                            value_transform="name_given_first_2",
                        ),
                        ExcelColumnMapping(
                            column_name="İsim",
                            field_name="hasta_soyad",
                            match_type="manual",
                            value_transform="name_family_first_2",
                        ),
                    ],
                ),
            )

        self.assertEqual(
            [(item.column_name, item.field_name, item.value_transform) for item in report.mapped_columns],
            [("İsim", "hasta_ad", "name_given_first_2"), ("İsim", "hasta_soyad", "name_family_first_2")],
        )
        fields = {item.field_name: item for item in report.results[0].merged_response.results}
        self.assertEqual(fields["hasta_ad"].final_value, "Ah")
        self.assertEqual(fields["hasta_soyad"].final_value, "Yı")

    def test_import_patient_excel_manual_mode_does_not_auto_map_unselected_columns(self) -> None:
        field_specs = {
            "hasta_cocuk_sayisi": FieldSpec(
                field_name="hasta_cocuk_sayisi",
                form_name="demografi",
                field_type="text",
                field_label="Çocuk sayısı",
            ),
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "Çocuk sayısı"])
            sheet.append(["77", 2])
            workbook.save(workbook_path)

            report = import_patient_excel(
                path=workbook_path,
                project_name="Demo",
                field_specs_by_name=field_specs,
                options=ExcelImportOptions(
                    use_llm_mapping=False,
                    automatic_mapping=False,
                    column_mappings=[],
                ),
            )

        self.assertEqual(report.mapped_columns, [])
        self.assertEqual(report.results[0].merged_response.results, [])

    def test_import_patient_excel_applies_field_post_processing(self) -> None:
        surname_spec = FieldSpec(
            field_name="hasta_soyad",
            form_name="hasta_bilgileri",
            field_type="text",
            field_label="Soyad",
        )
        surname_spec.post_processing = [["limit_output_length", 2]]
        field_specs = {"hasta_soyad": surname_spec}

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "Soyad"])
            sheet.append(["77", "Yılmaz"])
            workbook.save(workbook_path)

            report = import_patient_excel(
                path=workbook_path,
                project_name="Demo",
                field_specs_by_name=field_specs,
            )

        fields = {item.field_name: item for item in report.results[0].merged_response.results}
        self.assertEqual(fields["hasta_soyad"].final_value, "Yı")
        self.assertEqual(fields["hasta_soyad"].candidates[0]["value_normalized"], "Yı")

    def test_import_patient_excel_rejects_llm_name_mapping_from_operation_column(self) -> None:
        field_specs = {
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta Adı",
            ),
        }
        fake_provider = FakeProvider(
            '{"mappings":[{"column_name":"Radikal Prostatektomi Operasyon Tipi (??)",'
            '"field_name":"hasta_ad","confidence":0.99,"reason":"incorrect name match"}]}'
        )

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "Radikal Prostatektomi Operasyon Tipi (??)"])
            sheet.append(["77", "Robotik"])
            workbook.save(workbook_path)

            with patch("excel_import.create_provider", return_value=fake_provider):
                report = import_patient_excel(
                    path=workbook_path,
                    project_name="Demo",
                    field_specs_by_name=field_specs,
                    options=ExcelImportOptions(use_llm_mapping=True, llm_settings={"provider": "ollama", "model": "x"}),
                )

        self.assertEqual(report.mapped_columns, [])
        self.assertEqual(report.results[0].merged_response.results, [])

    def test_import_patient_excel_rejects_generic_name_column_for_middle_name(self) -> None:
        field_specs = {
            "hasta_gobek_adi": FieldSpec(
                field_name="hasta_gobek_adi",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Göbek Adı",
            ),
        }
        fake_provider = FakeProvider(
            '{"mappings":[{"column_name":"İsim","field_name":"hasta_gobek_adi",'
            '"confidence":0.94,"reason":"incorrect generic name match"}]}'
        )

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "İsim"])
            sheet.append(["77", "Ahmet Yılmaz"])
            workbook.save(workbook_path)

            with patch("excel_import.create_provider", return_value=fake_provider):
                report = import_patient_excel(
                    path=workbook_path,
                    project_name="Demo",
                    field_specs_by_name=field_specs,
                    options=ExcelImportOptions(use_llm_mapping=True, llm_settings={"provider": "ollama", "model": "x"}),
                )

        self.assertEqual(report.mapped_columns, [])
        self.assertEqual(report.results[0].merged_response.results, [])

    def test_import_patient_excel_rejects_llm_ln_count_for_child_count(self) -> None:
        field_specs = {
            "hasta_cocuk_sayisi": FieldSpec(
                field_name="hasta_cocuk_sayisi",
                form_name="demografi",
                field_type="text",
                field_label="Çocuk sayısı",
            ),
        }
        fake_provider = FakeProvider(
            '{"mappings":[{"column_name":"Çıkarılan LN sayısı",'
            '"field_name":"hasta_cocuk_sayisi","confidence":0.99,"reason":"incorrect count match"}]}'
        )

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "Çıkarılan LN sayısı"])
            sheet.append(["77", 14])
            workbook.save(workbook_path)

            with patch("excel_import.create_provider", return_value=fake_provider):
                report = import_patient_excel(
                    path=workbook_path,
                    project_name="Demo",
                    field_specs_by_name=field_specs,
                    options=ExcelImportOptions(use_llm_mapping=True, llm_settings={"provider": "ollama", "model": "x"}),
                )

        self.assertEqual(report.mapped_columns, [])
        self.assertEqual(report.results[0].merged_response.results, [])
        prompt_payload = fake_provider.calls[0]["messages"][1]["content"]
        self.assertNotIn("hasta_cocuk_sayisi", prompt_payload)

    def test_import_patient_excel_maps_ln_count_to_ln_field_without_llm(self) -> None:
        field_specs = {
            "rp_cikarilan_lenf_nodu_sayisi": FieldSpec(
                field_name="rp_cikarilan_lenf_nodu_sayisi",
                form_name="patoloji",
                field_type="text",
                field_label="Çıkarılan lenf nodu sayısı",
            ),
            "hasta_cocuk_sayisi": FieldSpec(
                field_name="hasta_cocuk_sayisi",
                form_name="demografi",
                field_type="text",
                field_label="Çocuk sayısı",
            ),
        }
        fake_provider = FakeProvider('{"mappings":[]}')

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "Çıkarılan LN sayısı"])
            sheet.append(["77", 14])
            workbook.save(workbook_path)

            with patch("excel_import.create_provider", return_value=fake_provider):
                report = import_patient_excel(
                    path=workbook_path,
                    project_name="Demo",
                    field_specs_by_name=field_specs,
                    options=ExcelImportOptions(use_llm_mapping=True, llm_settings={"provider": "ollama", "model": "x"}),
                )

        self.assertEqual(
            [(item.column_name, item.field_name, item.match_type) for item in report.mapped_columns],
            [("Çıkarılan LN sayısı", "rp_cikarilan_lenf_nodu_sayisi", "semantic_header")],
        )
        self.assertEqual(fake_provider.calls, [])
        fields = {item.field_name: item for item in report.results[0].merged_response.results}
        self.assertEqual(fields["rp_cikarilan_lenf_nodu_sayisi"].final_value, 14)

    def test_import_patient_excel_maps_common_synonym_without_llm(self) -> None:
        field_specs = {
            "hasta_kilo": FieldSpec(
                field_name="hasta_kilo",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Kilo (kg)",
                text_validation="integer",
            ),
        }
        fake_provider = FakeProvider(
            '{"mappings":[{"column_name":"Ağırlık","field_name":"wrong","confidence":1.0}]}'
        )

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "Ağırlık"])
            sheet.append(["77", 82])
            workbook.save(workbook_path)

            with patch("excel_import.create_provider", return_value=fake_provider):
                report = import_patient_excel(
                    path=workbook_path,
                    project_name="Demo",
                    field_specs_by_name=field_specs,
                    options=ExcelImportOptions(use_llm_mapping=True, llm_settings={"provider": "ollama", "model": "x"}),
                )

        self.assertEqual([(item.column_name, item.field_name, item.match_type) for item in report.mapped_columns], [("Ağırlık", "hasta_kilo", "semantic_header")])
        fields = {item.field_name: item for item in report.results[0].merged_response.results}
        self.assertEqual(fields["hasta_kilo"].final_value, 82)
        self.assertEqual(fake_provider.calls, [])

    def test_import_patient_excel_recovers_qwen_text_mapping(self) -> None:
        field_specs = {
            "hasta_kilo": FieldSpec(
                field_name="hasta_kilo",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Kilo (kg)",
                text_validation="integer",
            ),
        }
        fake_provider = FakeProvider("[hasta_kilo] ClinicMetric42 maps to field [field_name]")

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "ClinicMetric42"])
            sheet.append(["77", 82])
            workbook.save(workbook_path)

            with patch("excel_import.create_provider", return_value=fake_provider):
                report = import_patient_excel(
                    path=workbook_path,
                    project_name="Demo",
                    field_specs_by_name=field_specs,
                    options=ExcelImportOptions(use_llm_mapping=True, llm_settings={"provider": "ollama", "model": "x"}),
                )

        self.assertEqual(report.mapping_warnings, [])
        self.assertEqual([(item.column_name, item.field_name, item.match_type) for item in report.mapped_columns], [("ClinicMetric42", "hasta_kilo", "llm")])
        fields = {item.field_name: item for item in report.results[0].merged_response.results}
        self.assertEqual(fields["hasta_kilo"].final_value, 82)

    def test_import_patient_excel_does_not_fail_when_llm_mapping_is_unusable(self) -> None:
        field_specs = {
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta Adi",
            ),
        }
        fake_provider = FakeProvider("[unrelated_field] no useful mapping")

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "Hasta Adi", "Özel Kolon"])
            sheet.append(["77", "Ahmet", "x"])
            workbook.save(workbook_path)

            with patch("excel_import.create_provider", return_value=fake_provider):
                report = import_patient_excel(
                    path=workbook_path,
                    project_name="Demo",
                    field_specs_by_name=field_specs,
                    options=ExcelImportOptions(use_llm_mapping=True, llm_settings={"provider": "ollama", "model": "x"}),
                )

        self.assertTrue(any(item.startswith("llm_mapping_failed:") for item in report.mapping_warnings))
        self.assertEqual([mapping.field_name for mapping in report.mapped_columns], ["hasta_ad"])
        fields = {item.field_name: item for item in report.results[0].merged_response.results}
        self.assertEqual(fields["hasta_ad"].final_value, "Ahmet")

    def test_import_patient_excel_skips_openai_mapping_when_api_key_is_missing(self) -> None:
        field_specs = {
            "hasta_ad": FieldSpec(
                field_name="hasta_ad",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta Adi",
            ),
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "Hasta Adi", "ClinicMetric42"])
            sheet.append(["77", "Ahmet", "x"])
            workbook.save(workbook_path)

            with patch.dict(os.environ, {APP_HOME_ENV_VAR: tmp_dir}, clear=False):
                report = import_patient_excel(
                    path=workbook_path,
                    project_name="Demo",
                    field_specs_by_name=field_specs,
                    options=ExcelImportOptions(
                        use_llm_mapping=True,
                        llm_settings={"provider": "openai_compatible", "base_url": "https://example.com", "model": "x"},
                    ),
                )

        self.assertEqual(report.mapping_warnings, ["llm_mapping_skipped_missing_api_key"])
        self.assertEqual([mapping.field_name for mapping in report.mapped_columns], ["hasta_ad"])

    def test_import_patient_excel_handles_nan_dictionary_metadata(self) -> None:
        field_specs = {
            "hasta_kilo": FieldSpec(
                field_name="hasta_kilo",
                form_name="hasta_bilgileri",
                field_type="text",
                field_label="Hasta Kilo",
                field_note=float("nan"),
                text_validation=float("nan"),
            ),
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "patients.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["record_id", "Ağırlık"])
            sheet.append(["77", 82.5])
            workbook.save(workbook_path)

            report = import_patient_excel(
                path=workbook_path,
                project_name="Demo",
                field_specs_by_name=field_specs,
                options=ExcelImportOptions(use_llm_mapping=False),
            )

        self.assertEqual([mapping.field_name for mapping in report.mapped_columns], ["hasta_kilo"])
        fields = {item.field_name: item for item in report.results[0].merged_response.results}
        self.assertEqual(fields["hasta_kilo"].final_value, 82.5)


if __name__ == "__main__":
    unittest.main()
