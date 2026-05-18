from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
import re
import unicodedata
from typing import Any

from openpyxl import Workbook
from pydantic import BaseModel, ValidationError

from dictionary_parser import FieldSpec
from identity_registry_client import IdentityRegistryClient, IdentityRegistryError
from llm_provider import API_KEY_PROVIDER_NAMES, can_resolve_api_key, create_provider
from redcap_client import RedcapClient
from run_extraction import ExtractionFieldResult, normalize_json_response_text
from workspace_extraction import PatientExtractionResult
from workspace_review import serialize_review_value


class SubmissionValidationError(RuntimeError):
    pass


CHOICE_FIELD_TYPES = {"radio", "dropdown", "yesno", "checkbox"}
UNSUBMITTABLE_REVIEW_REASON = "choice_code_unmapped"
NUMERIC_ZERO_ABSENCE_KEYS = {"0", "no", "none", "absent", "nil", "negative", "yok", "yoktur", "hayir", "h", "y"}
NUMERIC_COUNT_TOKENS = {"adet", "count", "sayi", "sayisi", "total", "toplam"}
SUBMISSION_VALUE_REPAIR_BATCH_SIZE = 25


@dataclass
class UnsubmittableFieldIssue:
    result_index: int
    queue_label: str
    field_name: str
    field_label: str
    value: Any
    reason: str


@dataclass
class PatientSubmissionPlan:
    queue_label: str
    patient_mode: str
    action: str
    target_record_id: str | None
    tc_identity_no: str | None
    payload_rows: list[dict[str, Any]]
    warnings: list[str] = field(default_factory=list)

    @property
    def payload(self) -> dict[str, Any]:
        return self.payload_rows[0] if self.payload_rows else {}


class LLMSubmissionValueRepairItem(BaseModel):
    item_id: str
    normalized_value: Any = None
    value_code: Any = None
    confidence: float = 0.0
    reason: str | None = None


class LLMSubmissionValueRepairResponse(BaseModel):
    items: list[LLMSubmissionValueRepairItem]


@dataclass
class SubmissionValueRepairTask:
    item_id: str
    result_index: int
    field_name: str
    raw_value: Any
    reason: str


def export_review_results_to_excel(
    path: str | Path,
    results: list[PatientExtractionResult],
    field_specs_by_name: dict[str, FieldSpec],
) -> Path:
    workbook = Workbook()
    flat_sheet = workbook.active
    flat_sheet.title = "FlatDataRaw"
    coded_sheet = workbook.create_sheet("FlatDataCode")
    patient_sheet = workbook.create_sheet("Patients")
    patient_sheet.append(
        [
            "queue_label",
            "patient_mode",
            "identifier_type",
            "identifier_value",
            "approved",
            "field_count",
            "review_count",
        ]
    )
    field_sheet = workbook.create_sheet("Fields")
    field_sheet.append(
        [
            "queue_label",
            "patient_mode",
            "form_name",
            "field_name",
            "field_label",
            "field_type",
            "status",
            "final_value",
            "final_value_code",
            "confidence",
            "needs_review",
            "selection_reason",
            "review_reasons",
        ]
    )
    result_field_names = [
        field_result.field_name
        for result in results
        for field_result in result.merged_response.results
    ]
    all_field_names = order_field_names(result_field_names, field_specs_by_name)
    coded_field_names = [
        field_name
        for field_name in all_field_names
        if any(
            field_result.field_name == field_name and field_result.final_value_code not in {None, ""}
            for result in results
            for field_result in result.merged_response.results
        )
    ]
    base_header = [
        "queue_label",
        "patient_mode",
        "identifier_type",
        "identifier_value",
        "approved",
    ]
    flat_sheet.append(base_header + all_field_names)
    coded_sheet.append(base_header + coded_field_names)

    for result in results:
        review_count = sum(1 for item in result.merged_response.results if item.needs_review)
        patient_sheet.append(
            [
                result.queue_label,
                result.patient_mode,
                result.identifier_type,
                result.identifier_value,
                "yes" if result.approved else "no",
                len(result.merged_response.results),
                review_count,
            ]
        )
        for field_result in result.merged_response.results:
            field_spec = field_specs_by_name.get(field_result.field_name)
            field_sheet.append(
                [
                    result.queue_label,
                    result.patient_mode,
                    field_result.form_name,
                    field_result.field_name,
                    field_spec.field_label if field_spec else "",
                    field_spec.field_type if field_spec else "",
                    field_result.status,
                    serialize_review_value(field_result.final_value),
                    field_result.final_value_code,
                    field_result.confidence,
                    "yes" if field_result.needs_review else "no",
                    field_result.selection_reason,
                    " | ".join(field_result.review_reasons or []),
                ]
            )
        field_map = {item.field_name: item for item in result.merged_response.results}
        base_row: list[Any] = [
            result.queue_label,
            result.patient_mode,
            result.identifier_type,
            result.identifier_value,
            "yes" if result.approved else "no",
        ]
        flat_row = list(base_row)
        coded_row = list(base_row)
        for field_name in all_field_names:
            field_result = field_map.get(field_name)
            flat_row.append(serialize_review_value(field_result.final_value) if field_result else "")
        for field_name in coded_field_names:
            field_result = field_map.get(field_name)
            coded_row.append(field_result.final_value_code if field_result else "")
        flat_sheet.append(flat_row)
        coded_sheet.append(coded_row)

    destination = Path(path).expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination


def build_field_specs_by_name(
    grouped_fields: dict[str, list[FieldSpec]],
    append_fields: list[dict[str, Any]] | None = None,
    dictionary_legend: dict[str, str] | None = None,
) -> dict[str, FieldSpec]:
    field_specs: dict[str, FieldSpec] = {}
    for fields in grouped_fields.values():
        for field in fields:
            field_specs[field.field_name] = field
    legend = dictionary_legend or {}
    source_field_name = legend.get("field_name", "Variable / Field Name")
    source_form_name = legend.get("form_name", "Form Name")
    source_field_type = legend.get("field_type", "Field Type")
    source_field_label = legend.get("field_label", "Field Label")
    source_choices = legend.get("choices", "Choices, Calculations, OR Slider Labels")
    source_field_note = legend.get("field_note", "Field Note")
    source_text_validation = legend.get("text_validation", "Text Validation Type OR Show Slider Number")
    source_text_validation_min = legend.get("text_validation_min", "Text Validation Min")
    source_text_validation_max = legend.get("text_validation_max", "Text Validation Max")
    source_field_annotation = legend.get("field_annotation", "Field Annotation")
    for raw in append_fields or []:
        field_name = str(raw.get(source_field_name, "")).strip()
        if not field_name:
            continue
        if field_name in field_specs:
            continue
        field_specs[field_name] = FieldSpec(
            field_name=field_name,
            form_name=str(raw.get(source_form_name, "")).strip(),
            field_type=str(raw.get(source_field_type, "text") or "text"),
            field_label=str(raw.get(source_field_label, field_name) or field_name),
            choices=str(raw.get(source_choices, "") or ""),
            field_note=str(raw.get(source_field_note, "") or ""),
            text_validation=str(raw.get(source_text_validation, "") or ""),
            text_validation_min=str(raw.get(source_text_validation_min, "") or ""),
            text_validation_max=str(raw.get(source_text_validation_max, "") or ""),
            field_annotation=str(raw.get(source_field_annotation, "") or ""),
        )
    return field_specs


def build_append_field_name_set(
    append_fields: list[dict[str, Any]] | None = None,
    dictionary_legend: dict[str, str] | None = None,
) -> set[str]:
    legend = dictionary_legend or {}
    source_field_name = legend.get("field_name", "Variable / Field Name")
    names: set[str] = set()
    for raw in append_fields or []:
        field_name = str(raw.get(source_field_name, "")).strip()
        if field_name:
            names.add(field_name)
    return names


def order_field_names(field_names: list[str], field_specs_by_name: dict[str, FieldSpec]) -> list[str]:
    seen: set[str] = set()
    ordered: list[str] = []

    for field_name in field_specs_by_name.keys():
        if field_name in seen:
            continue
        if field_name not in field_names:
            continue
        seen.add(field_name)
        ordered.append(field_name)

    for field_name in field_names:
        if field_name in seen:
            continue
        seen.add(field_name)
        ordered.append(field_name)

    return ordered


def resolve_patient_submission_plan(
    *,
    result: PatientExtractionResult,
    field_specs_by_name: dict[str, FieldSpec],
    append_field_names: set[str] | None,
    repeating_forms: set[str] | None,
    repeating_events: set[str] | None,
    form_event_map: dict[str, list[str]] | None,
    identity_client: IdentityRegistryClient | None,
    manual_tc_identity_no: str | None = None,
) -> PatientSubmissionPlan:
    tc_identity_no = (
        extract_tc_identity_no(result)
        or (
            normalize_tc_identity_no(result.identifier_value)
            if result.identifier_type == "tc_kimlik_no"
            else None
        )
        or normalize_tc_identity_no(manual_tc_identity_no)
    )
    record_id = extract_record_id(result) or (
        (result.identifier_value or "").strip() if result.identifier_type == "record_id" else None
    )

    if result.patient_mode == "auto":
        if record_id:
            payload_rows, warnings = build_redcap_record_payload_rows(
                result=result,
                target_record_id=record_id,
                field_specs_by_name=field_specs_by_name,
                excluded_field_names=append_field_names,
                repeating_forms=repeating_forms,
                repeating_events=repeating_events,
                form_event_map=form_event_map,
            )
            return PatientSubmissionPlan(
                queue_label=result.queue_label,
                patient_mode=result.patient_mode,
                action="update",
                target_record_id=record_id,
                tc_identity_no=tc_identity_no,
                payload_rows=payload_rows,
                warnings=warnings,
            )
        if not tc_identity_no:
            raise SubmissionValidationError("missing_tc_identity")
        if identity_client is None:
            raise SubmissionValidationError("identity_registry_required")
        existing_record_id = identity_client.lookup_record_id(tc_identity_no)
        if existing_record_id:
            payload_rows, warnings = build_redcap_record_payload_rows(
                result=result,
                target_record_id=existing_record_id,
                field_specs_by_name=field_specs_by_name,
                excluded_field_names=append_field_names,
                repeating_forms=repeating_forms,
                repeating_events=repeating_events,
                form_event_map=form_event_map,
            )
            return PatientSubmissionPlan(
                queue_label=result.queue_label,
                patient_mode=result.patient_mode,
                action="update",
                target_record_id=existing_record_id,
                tc_identity_no=tc_identity_no,
                payload_rows=payload_rows,
                warnings=warnings,
            )
        create_result = identity_client.create_record_by_tc(tc_identity_no)
        if not create_result.record_id:
            raise SubmissionValidationError("identity_registry_create_failed")
        payload_rows, warnings = build_redcap_record_payload_rows(
            result=result,
            target_record_id=create_result.record_id,
            field_specs_by_name=field_specs_by_name,
            excluded_field_names=append_field_names,
            repeating_forms=repeating_forms,
            repeating_events=repeating_events,
            form_event_map=form_event_map,
        )
        return PatientSubmissionPlan(
            queue_label=result.queue_label,
            patient_mode=result.patient_mode,
            action="create" if create_result.created else "update",
            target_record_id=create_result.record_id,
            tc_identity_no=tc_identity_no,
            payload_rows=payload_rows,
            warnings=warnings,
        )

    if result.patient_mode == "new":
        if not tc_identity_no:
            raise SubmissionValidationError("missing_tc_identity")
        if identity_client is None:
            raise SubmissionValidationError("identity_registry_required")
        existing_record_id = identity_client.lookup_record_id(tc_identity_no)
        if existing_record_id:
            raise SubmissionValidationError(f"duplicate_tc_identity:{existing_record_id}")
        create_result = identity_client.create_record_by_tc(tc_identity_no)
        if not create_result.record_id:
            raise SubmissionValidationError("identity_registry_create_failed")
        if not create_result.created:
            raise SubmissionValidationError(f"duplicate_tc_identity:{create_result.record_id}")
        payload_rows, warnings = build_redcap_record_payload_rows(
            result=result,
            target_record_id=create_result.record_id,
            field_specs_by_name=field_specs_by_name,
            excluded_field_names=append_field_names,
            repeating_forms=repeating_forms,
            repeating_events=repeating_events,
            form_event_map=form_event_map,
        )
        return PatientSubmissionPlan(
            queue_label=result.queue_label,
            patient_mode=result.patient_mode,
            action="create",
            target_record_id=create_result.record_id,
            tc_identity_no=tc_identity_no,
            payload_rows=payload_rows,
            warnings=warnings,
        )

    if result.identifier_type == "record_id":
        if not record_id:
            raise SubmissionValidationError("missing_record_id")
        payload_rows, warnings = build_redcap_record_payload_rows(
            result=result,
            target_record_id=record_id,
            field_specs_by_name=field_specs_by_name,
            excluded_field_names=append_field_names,
            repeating_forms=repeating_forms,
            repeating_events=repeating_events,
            form_event_map=form_event_map,
        )
        return PatientSubmissionPlan(
            queue_label=result.queue_label,
            patient_mode=result.patient_mode,
            action="update",
            target_record_id=record_id,
            tc_identity_no=None,
            payload_rows=payload_rows,
            warnings=warnings,
        )

    if result.identifier_type == "tc_kimlik_no":
        lookup_tc = normalize_tc_identity_no(result.identifier_value)
        if not lookup_tc:
            raise SubmissionValidationError("missing_tc_identity")
        if identity_client is None:
            raise SubmissionValidationError("identity_registry_required")
        record_id = identity_client.lookup_record_id(lookup_tc)
        if not record_id:
            raise SubmissionValidationError("record_id_not_found_for_tc")
        payload_rows, warnings = build_redcap_record_payload_rows(
            result=result,
            target_record_id=record_id,
            field_specs_by_name=field_specs_by_name,
            excluded_field_names=append_field_names,
            repeating_forms=repeating_forms,
            repeating_events=repeating_events,
            form_event_map=form_event_map,
        )
        return PatientSubmissionPlan(
            queue_label=result.queue_label,
            patient_mode=result.patient_mode,
            action="update",
            target_record_id=record_id,
            tc_identity_no=lookup_tc,
            payload_rows=payload_rows,
            warnings=warnings,
        )

    raise SubmissionValidationError("unsupported_patient_identifier")


def submit_patient_plan(
    *,
    plan: PatientSubmissionPlan,
    redcap_client: RedcapClient,
    identity_client: IdentityRegistryClient | None,
) -> str:
    force_auto_number = False
    if not plan.payload_rows:
        if plan.target_record_id:
            return plan.target_record_id
        raise SubmissionValidationError("no_submittable_fields")
    returned_ids = redcap_client.import_records(
        plan.payload_rows,
        force_auto_number=force_auto_number,
        return_content="ids",
    )
    if plan.action == "update":
        return plan.target_record_id or ""
    if plan.target_record_id:
        return plan.target_record_id

    if not returned_ids:
        raise SubmissionValidationError("redcap_auto_id_not_returned")
    new_record_id = str(returned_ids[0])
    return new_record_id


def build_redcap_record_payload(
    *,
    result: PatientExtractionResult,
    target_record_id: str | None,
    field_specs_by_name: dict[str, FieldSpec],
    excluded_field_names: set[str] | None = None,
) -> tuple[dict[str, Any], list[str]]:
    payload: dict[str, Any] = {}
    warnings: list[str] = []
    excluded = excluded_field_names or set()

    if target_record_id:
        payload["record_id"] = target_record_id

    for field_result in result.merged_response.results:
        if field_result.final_value is None or field_result.status == "not_found":
            continue
        if field_result.field_name in excluded:
            continue
        field_spec = field_specs_by_name.get(field_result.field_name)
        field_type = field_spec.field_type if field_spec else "text"

        if field_type in {"radio", "dropdown", "yesno"}:
            code = single_choice_code_for_submission(field_result, field_spec)
            if code is None:
                warnings.append(f"{field_result.field_name}: skipped_choice_code_unmapped")
                continue
            payload[field_result.field_name] = code
            continue

        if field_type == "checkbox":
            codes = checkbox_codes_for_submission(field_result, field_spec)
            if not codes:
                warnings.append(f"{field_result.field_name}: unsupported_checkbox_value")
                continue
            for code in codes:
                payload[f"{field_result.field_name}___{code}"] = "1"
            continue

        if isinstance(field_result.final_value, list):
            warnings.append(f"{field_result.field_name}: skipped_multiple_value")
            continue

        normalized_value, value_warning = normalize_field_value_for_submission(
            field_result.field_name,
            field_result.final_value,
            field_spec,
        )
        if value_warning:
            warnings.append(value_warning)
            continue
        payload[field_result.field_name] = normalized_value

    return payload, warnings


def build_redcap_record_payload_rows(
    *,
    result: PatientExtractionResult,
    target_record_id: str | None,
    field_specs_by_name: dict[str, FieldSpec],
    excluded_field_names: set[str] | None = None,
    repeating_forms: set[str] | None = None,
    repeating_events: set[str] | None = None,
    form_event_map: dict[str, list[str]] | None = None,
) -> tuple[list[dict[str, Any]], list[str]]:
    repeating_form_names = repeating_forms or set()
    repeating_event_names = repeating_events or set()
    form_events = form_event_map or {}
    warnings: list[str] = []
    base_rows: dict[str | None, dict[str, Any]] = {}
    repeating_rows: dict[tuple[str | None, str], dict[str, Any]] = {}
    excluded = excluded_field_names or set()

    for field_result in result.merged_response.results:
        if field_result.final_value is None or field_result.status == "not_found":
            continue
        if field_result.field_name in excluded:
            continue
        field_spec = field_specs_by_name.get(field_result.field_name)
        form_name = field_spec.form_name if field_spec else field_result.form_name
        event_name = first_form_event_name(form_name, form_events)
        row_key = event_name
        if form_name in repeating_form_names:
            row = repeating_rows.setdefault(
                (event_name, form_name or ""),
                {
                    "record_id": target_record_id,
                    "redcap_repeat_instrument": form_name,
                    "redcap_repeat_instance": "new",
                },
            )
            if event_name:
                row["redcap_event_name"] = event_name
        else:
            row = base_rows.setdefault(
                row_key,
                {
                    "record_id": target_record_id,
                },
            )
            if event_name:
                row["redcap_event_name"] = event_name
            if event_name in repeating_event_names:
                row["redcap_repeat_instance"] = "new"

        field_type = field_spec.field_type if field_spec else "text"
        if field_type in {"radio", "dropdown", "yesno"}:
            code = single_choice_code_for_submission(field_result, field_spec)
            if code is None:
                warnings.append(f"{field_result.field_name}: skipped_choice_code_unmapped")
                continue
            row[field_result.field_name] = code
            continue
        if field_type == "checkbox":
            codes = checkbox_codes_for_submission(field_result, field_spec)
            if not codes:
                warnings.append(f"{field_result.field_name}: unsupported_checkbox_value")
                continue
            for code in codes:
                row[f"{field_result.field_name}___{code}"] = "1"
            continue
        if isinstance(field_result.final_value, list):
            warnings.append(f"{field_result.field_name}: skipped_multiple_value")
            continue
        normalized_value, value_warning = normalize_field_value_for_submission(
            field_result.field_name,
            field_result.final_value,
            field_spec,
        )
        if value_warning:
            warnings.append(value_warning)
            continue
        row[field_result.field_name] = normalized_value

    rows: list[dict[str, Any]] = []
    rows.extend(
        row
        for row in base_rows.values()
        if any(key not in {"record_id", "redcap_event_name", "redcap_repeat_instance"} for key in row.keys())
    )
    rows.extend(
        row
        for row in repeating_rows.values()
        if any(
            key not in {"record_id", "redcap_event_name", "redcap_repeat_instrument", "redcap_repeat_instance"}
            for key in row.keys()
        )
    )
    return rows, warnings


def first_form_event_name(form_name: str | None, form_event_map: dict[str, list[str]]) -> str | None:
    if not form_name:
        return None
    events = form_event_map.get(form_name) or []
    if not events:
        return None
    return str(events[0]).strip() or None


def normalize_scalar_value(value: Any) -> Any:
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float)):
        return value
    if value is None:
        return None
    return str(value)


def normalize_field_value_for_submission(
    field_name: str,
    value: Any,
    field_spec: FieldSpec | None,
) -> tuple[Any, str | None]:
    if is_redcap_date_field(field_spec):
        normalized_date = normalize_redcap_date_value(value)
        if normalized_date:
            return normalized_date, None
        return None, f"{field_name}: skipped_invalid_date_format"
    if is_redcap_numeric_field(field_spec):
        normalized_number, number_warning = normalize_redcap_numeric_value(value, field_spec)
        if number_warning:
            return None, f"{field_name}: {number_warning}"
        return normalized_number, None
    return normalize_scalar_value(value), None


def is_redcap_date_field(field_spec: FieldSpec | None) -> bool:
    return bool(field_spec and str(field_spec.text_validation or "").strip().lower().startswith("date"))


def redcap_text_validation(field_spec: FieldSpec | None) -> str:
    return str(field_spec.text_validation or "").strip().lower() if field_spec else ""


def is_redcap_integer_field(field_spec: FieldSpec | None) -> bool:
    return redcap_text_validation(field_spec) == "integer"


def is_redcap_number_field(field_spec: FieldSpec | None) -> bool:
    validation = redcap_text_validation(field_spec)
    return validation in {"number", "float"} or validation.startswith("number")


def is_redcap_numeric_field(field_spec: FieldSpec | None) -> bool:
    return is_redcap_integer_field(field_spec) or is_redcap_number_field(field_spec)


def normalize_redcap_numeric_value(value: Any, field_spec: FieldSpec | None) -> tuple[int | float | None, str | None]:
    if numeric_absence_value_should_be_zero(value, field_spec):
        return 0, None
    numeric_value = parse_submission_float(value)
    if numeric_value is None:
        return None, "skipped_invalid_number_format"

    if is_redcap_integer_field(field_spec):
        if not numeric_value.is_integer():
            return None, "skipped_invalid_number_format"
        normalized_value: int | float = int(numeric_value)
    else:
        normalized_value = int(numeric_value) if numeric_value.is_integer() else numeric_value

    min_value = parse_submission_float(getattr(field_spec, "text_validation_min", None))
    max_value = parse_submission_float(getattr(field_spec, "text_validation_max", None))
    if min_value is not None and float(normalized_value) < min_value:
        return None, "skipped_number_out_of_range"
    if max_value is not None and float(normalized_value) > max_value:
        return None, "skipped_number_out_of_range"
    return normalized_value, None


def numeric_absence_value_should_be_zero(value: Any, field_spec: FieldSpec | None) -> bool:
    if not is_redcap_numeric_field(field_spec):
        return False
    if not numeric_field_represents_count(field_spec):
        return False
    return normalize_choice_match_key(value) in NUMERIC_ZERO_ABSENCE_KEYS


def numeric_field_represents_count(field_spec: FieldSpec | None) -> bool:
    if field_spec is None:
        return False
    text = normalize_submission_search_text(
        " ".join(
            [
                field_spec.field_name,
                field_spec.form_name,
                field_spec.field_label,
                field_spec.field_note or "",
            ]
        )
    )
    tokens = set(text.split())
    return bool(tokens & NUMERIC_COUNT_TOKENS)


def normalize_submission_search_text(value: Any) -> str:
    text = str(value or "").strip().casefold().replace("ı", "i")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def parse_submission_float(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        numeric_value = float(value)
        return numeric_value if numeric_value == numeric_value and numeric_value not in {float("inf"), float("-inf")} else None
    text = str(value).strip()
    if not text:
        return None
    normalized = text.replace(" ", "")
    if not re.fullmatch(r"[-+]?\d+(?:[.,]\d+)?", normalized):
        return None
    try:
        return float(normalized.replace(",", "."))
    except ValueError:
        return None


def normalize_redcap_date_value(value: Any) -> str | None:
    if is_blank_submission_value(value):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return text
    for pattern in ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%y", "%d/%m/%y", "%d-%m-%y"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    return None


def single_choice_code_for_submission(
    result: ExtractionFieldResult,
    field_spec: FieldSpec | None,
) -> str | None:
    if field_spec and field_spec.field_type == "yesno":
        return normalize_yesno_code_for_submission(result.final_value_code) or normalize_yesno_code_for_submission(
            result.final_value
        )
    if field_spec is None or field_spec.field_type not in {"radio", "dropdown"}:
        return normalize_scalar_value(result.final_value_code) if result.final_value_code not in {None, ""} else None
    code = match_choice_code(result.final_value_code, field_spec)
    if code is not None:
        return code
    return match_choice_code(result.final_value, field_spec)


def checkbox_codes_for_submission(result: ExtractionFieldResult, field_spec: FieldSpec | None) -> list[str]:
    if result.needs_review and UNSUBMITTABLE_REVIEW_REASON in set(result.review_reasons or []):
        return []
    code_values = normalize_checkbox_codes(result)
    if field_spec is not None and field_spec.choices_options:
        valid_codes = {str(option.code).strip() for option in field_spec.choices_options}
        code_values = [code for code in code_values if code in valid_codes]
        if code_values:
            return code_values
        return match_checkbox_value_codes(result.final_value, field_spec)
    return code_values


def match_checkbox_value_codes(value: Any, field_spec: FieldSpec) -> list[str]:
    parts = value if isinstance(value, list) else split_checkbox_submission_value(value)
    codes: list[str] = []
    for part in parts:
        code = match_choice_code(part, field_spec)
        if code is None:
            return []
        codes.append(code)
    return codes


def split_checkbox_submission_value(value: Any) -> list[Any]:
    if is_blank_submission_value(value):
        return []
    if isinstance(value, list):
        return [item for item in value if not is_blank_submission_value(item)]
    text = str(value).strip()
    if not text:
        return []
    separators = r"[;|\n]"
    if not re.search(separators, text) and "," in text:
        separators = r","
    return [part.strip() for part in re.split(separators, text) if part.strip()]


def match_choice_code(value: Any, field_spec: FieldSpec) -> str | None:
    if is_blank_submission_value(value):
        return None
    value_text = str(value).strip()
    normalized_value = normalize_choice_match_key(value_text)
    for option in field_spec.choices_options:
        code = str(option.code).strip()
        if value_text == code or normalized_value == normalize_choice_match_key(code):
            return code
    for option in field_spec.choices_options:
        if normalized_value == normalize_choice_match_key(option.label):
            return str(option.code).strip()
    return None


def normalize_choice_match_key(value: Any) -> str:
    text = str(value).strip().casefold().replace("ı", "i")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", text)


def normalize_yesno_code_for_submission(value: Any) -> str | None:
    normalized = normalize_choice_match_key(value)
    if normalized in {"1", "true", "yes", "y", "evet", "e", "var"}:
        return "1"
    if normalized in {"0", "false", "no", "n", "hayir", "h", "yok"}:
        return "0"
    return None


def repair_unsubmittable_fields_for_submission(
    results: list[PatientExtractionResult],
    field_specs_by_name: dict[str, FieldSpec],
    *,
    llm_settings: dict[str, Any] | None = None,
    progress_callback: Any | None = None,
) -> tuple[int, list[str]]:
    repaired_count = apply_procedural_submission_value_repairs(results, field_specs_by_name)
    warnings: list[str] = []

    remaining_issues = collect_unsubmittable_review_fields(results, field_specs_by_name)
    if progress_callback is not None:
        progress_callback(0, len(remaining_issues), "submission_value_repair")
    if not remaining_issues:
        return repaired_count, warnings

    if not llm_settings:
        return repaired_count, warnings
    if not can_call_submission_llm(llm_settings):
        warnings.append("llm_submission_value_repair_skipped_missing_api_key")
        return repaired_count, warnings

    try:
        repaired_count += repair_submission_values_with_llm(
            results=results,
            field_specs_by_name=field_specs_by_name,
            issues=remaining_issues,
            llm_settings=llm_settings,
            progress_callback=progress_callback,
        )
    except ValueError as exc:
        warnings.append(f"llm_submission_value_repair_failed:{str(exc)[:300]}")
    return repaired_count, warnings


def apply_procedural_submission_value_repairs(
    results: list[PatientExtractionResult],
    field_specs_by_name: dict[str, FieldSpec],
) -> int:
    repaired_count = 0
    for result in results:
        for field_result in result.merged_response.results:
            field_spec = field_specs_by_name.get(field_result.field_name)
            if field_spec is None:
                continue
            if numeric_absence_value_should_be_zero(field_result.final_value, field_spec):
                apply_submission_repair_to_result(
                    field_result=field_result,
                    normalized_value=0,
                    value_code=None,
                    confidence=1.0,
                    reason="Sayısal adet alanında yokluk bildirimi 0 olarak normalize edildi.",
                )
                repaired_count += 1
    return repaired_count


def can_call_submission_llm(settings: dict[str, Any]) -> bool:
    provider_name = str(settings.get("provider", "ollama")).strip().lower()
    if provider_name == "ollama":
        return True
    if provider_name not in API_KEY_PROVIDER_NAMES:
        return True
    return can_resolve_api_key(settings)


def repair_submission_values_with_llm(
    *,
    results: list[PatientExtractionResult],
    field_specs_by_name: dict[str, FieldSpec],
    issues: list[UnsubmittableFieldIssue],
    llm_settings: dict[str, Any],
    progress_callback: Any | None = None,
) -> int:
    tasks = build_submission_value_repair_tasks(issues)
    if not tasks:
        return 0

    settings = build_submission_value_repair_llm_settings(llm_settings)
    provider = create_provider(settings)
    schema = LLMSubmissionValueRepairResponse.model_json_schema()
    repaired_items: dict[str, LLMSubmissionValueRepairItem] = {}
    for offset in range(0, len(tasks), SUBMISSION_VALUE_REPAIR_BATCH_SIZE):
        batch = tasks[offset : offset + SUBMISSION_VALUE_REPAIR_BATCH_SIZE]
        messages = build_submission_value_repair_messages(batch, field_specs_by_name)
        response = provider.generate(messages, schema, settings)
        content = response.content
        if not content:
            raise ValueError("LLM submission value repair returned an empty response.")
        parsed = parse_llm_submission_value_repair_response(content)
        for item in parsed.items:
            if float(item.confidence or 0.0) >= 0.65:
                repaired_items[item.item_id] = item
        if progress_callback is not None:
            progress_callback(
                min(offset + len(batch), len(tasks)),
                len(tasks),
                "submission_value_repair",
            )
    return apply_llm_submission_value_repairs(
        results=results,
        field_specs_by_name=field_specs_by_name,
        tasks=tasks,
        repaired_items=repaired_items,
    )


def build_submission_value_repair_tasks(issues: list[UnsubmittableFieldIssue]) -> list[SubmissionValueRepairTask]:
    tasks: list[SubmissionValueRepairTask] = []
    seen: set[tuple[int, str, str]] = set()
    for issue in issues:
        key = (issue.result_index, issue.field_name, serialize_review_value(issue.value))
        if key in seen:
            continue
        seen.add(key)
        tasks.append(
            SubmissionValueRepairTask(
                item_id=f"r{len(tasks) + 1}",
                result_index=issue.result_index,
                field_name=issue.field_name,
                raw_value=issue.value,
                reason=issue.reason,
            )
        )
    return tasks


def build_submission_value_repair_llm_settings(llm_settings: dict[str, Any]) -> dict[str, Any]:
    settings = dict(llm_settings)
    settings["temperature"] = 0
    settings["think"] = False
    try:
        current_max_tokens = int(settings.get("max_tokens", 2048) or 2048)
    except (TypeError, ValueError):
        current_max_tokens = 2048
    settings["max_tokens"] = min(max(current_max_tokens, 768), 2048)
    return settings


def build_submission_value_repair_messages(
    tasks: list[SubmissionValueRepairTask],
    field_specs_by_name: dict[str, FieldSpec],
) -> list[dict[str, str]]:
    items: list[dict[str, Any]] = []
    for task in tasks:
        field_spec = field_specs_by_name[task.field_name]
        items.append(
            {
                "item_id": task.item_id,
                "field_name": task.field_name,
                "form_name": field_spec.form_name,
                "field_label": field_spec.field_label,
                "field_type": field_spec.field_type,
                "text_validation": field_spec.text_validation,
                "text_validation_min": field_spec.text_validation_min,
                "text_validation_max": field_spec.text_validation_max,
                "choices": [
                    {"code": choice.code, "label": choice.label}
                    for choice in field_spec.choices_options
                ],
                "raw_value": task.raw_value,
                "problem": task.reason,
            }
        )
    return [
        {
            "role": "system",
            "content": (
                "You repair values before a REDCap import. Return valid JSON only. "
                "Be conservative and do not invent clinical facts. If a value cannot be safely converted, return nulls."
            ),
        },
        {
            "role": "user",
            "content": json.dumps(
                {
                    "task": "Convert invalid values into REDCap-compatible values.",
                    "rules": [
                        "Return one item per input item_id.",
                        "For integer/number fields, normalized_value must be numeric.",
                        "For count fields, absence words such as yok/no/none may be converted to 0 when clinically appropriate.",
                        "For date fields, normalized_value must be YYYY-MM-DD.",
                        "For radio/dropdown/yesno fields, value_code must be one of the provided choice codes.",
                        "For checkbox fields, value_code may be an array of provided choice codes.",
                        "If the value is likely from a wrong Excel column or cannot be converted, return normalized_value null and value_code null.",
                    ],
                    "items": items,
                    "response_schema": {
                        "items": [
                            {
                                "item_id": "same id from input",
                                "normalized_value": "string | number | boolean | array | null",
                                "value_code": "string | array | null",
                                "confidence": 0.0,
                                "reason": "short reason",
                            }
                        ]
                    },
                },
                ensure_ascii=False,
            ),
        },
    ]


def parse_llm_submission_value_repair_response(content: str) -> LLMSubmissionValueRepairResponse:
    normalized = normalize_json_response_text(content)
    try:
        return LLMSubmissionValueRepairResponse.model_validate_json(normalized)
    except ValidationError as json_exc:
        first_object = normalized.find("{")
        last_object = normalized.rfind("}")
        candidates = [normalized]
        if first_object != -1 and last_object > first_object:
            candidates.append(normalized[first_object : last_object + 1].strip())
        first_array = normalized.find("[")
        last_array = normalized.rfind("]")
        if first_array != -1 and last_array > first_array:
            candidates.append(normalized[first_array : last_array + 1].strip())
        for candidate in candidates:
            try:
                raw = json.loads(candidate)
            except json.JSONDecodeError:
                continue
            if isinstance(raw, dict) and isinstance(raw.get("items"), list):
                return LLMSubmissionValueRepairResponse.model_validate(raw)
            if isinstance(raw, list):
                return LLMSubmissionValueRepairResponse.model_validate({"items": raw})
        raise ValueError(f"invalid_json:{json_exc}") from json_exc


def apply_llm_submission_value_repairs(
    *,
    results: list[PatientExtractionResult],
    field_specs_by_name: dict[str, FieldSpec],
    tasks: list[SubmissionValueRepairTask],
    repaired_items: dict[str, LLMSubmissionValueRepairItem],
) -> int:
    repaired_count = 0
    for task in tasks:
        item = repaired_items.get(task.item_id)
        if item is None:
            continue
        if task.result_index < 0 or task.result_index >= len(results):
            continue
        field_result = find_submission_field_result(
            results[task.result_index],
            task.field_name,
            task.raw_value,
        )
        field_spec = field_specs_by_name.get(task.field_name)
        if field_result is None or field_spec is None:
            continue
        if apply_llm_submission_repair_item(field_result, field_spec, item):
            repaired_count += 1
    return repaired_count


def find_submission_field_result(
    result: PatientExtractionResult,
    field_name: str,
    raw_value: Any,
) -> ExtractionFieldResult | None:
    raw_key = serialize_review_value(raw_value)
    for field_result in result.merged_response.results:
        if field_result.field_name != field_name:
            continue
        if serialize_review_value(field_result.final_value) == raw_key:
            return field_result
    return None


def apply_llm_submission_repair_item(
    field_result: ExtractionFieldResult,
    field_spec: FieldSpec,
    item: LLMSubmissionValueRepairItem,
) -> bool:
    candidate_value = item.normalized_value
    candidate_code = item.value_code
    if field_spec.field_type in {"radio", "dropdown", "yesno"}:
        code = match_choice_code(candidate_code, field_spec) or match_choice_code(candidate_value, field_spec)
        if code is None:
            return False
        label = choice_label_for_code(code, field_spec)
        apply_submission_repair_to_result(field_result, label or candidate_value, code, float(item.confidence), item.reason)
        return True
    if field_spec.field_type == "checkbox":
        candidate = candidate_code if candidate_code not in {None, ""} else candidate_value
        codes = match_checkbox_value_codes(candidate, field_spec)
        if not codes:
            return False
        labels = [choice_label_for_code(code, field_spec) or code for code in codes]
        apply_submission_repair_to_result(
            field_result,
            labels,
            json.dumps(codes, ensure_ascii=False),
            float(item.confidence),
            item.reason,
        )
        return True
    normalized_value, warning = normalize_field_value_for_submission(
        field_result.field_name,
        candidate_value,
        field_spec,
    )
    if warning:
        return False
    apply_submission_repair_to_result(
        field_result,
        normalized_value,
        None,
        float(item.confidence),
        item.reason,
    )
    return True


def choice_label_for_code(code: Any, field_spec: FieldSpec) -> str | None:
    matched_code = match_choice_code(code, field_spec)
    if matched_code is None:
        return None
    for option in field_spec.choices_options:
        if str(option.code).strip() == matched_code:
            return option.label
    return None


def apply_submission_repair_to_result(
    field_result: ExtractionFieldResult,
    normalized_value: Any,
    value_code: Any,
    confidence: float,
    reason: str | None,
) -> None:
    original_value = field_result.final_value
    field_result.final_value = normalized_value
    field_result.final_value_code = value_code
    field_result.status = "found"
    field_result.needs_review = False
    field_result.review_reasons = []
    field_result.confidence = max(0.0, min(confidence, 1.0))
    field_result.selection_reason = "submission_value_repair"
    field_result.candidates = [
        {
            "value_raw": serialize_review_value(original_value),
            "value_normalized": normalized_value,
            "value_code": value_code,
            "confidence": field_result.confidence,
            "evidence": reason or "submission value repair",
            "context_date": None,
            "context_label": "submission_value_repair",
        },
        *(field_result.candidates or []),
    ]


def collect_unsubmittable_review_fields(
    results: list[PatientExtractionResult],
    field_specs_by_name: dict[str, FieldSpec],
) -> list[UnsubmittableFieldIssue]:
    issues: list[UnsubmittableFieldIssue] = []
    for result_index, result in enumerate(results):
        for field_result in result.merged_response.results:
            if is_blank_submission_value(field_result.final_value) or field_result.status == "not_found":
                continue
            field_spec = field_specs_by_name.get(field_result.field_name)
            if field_spec is None:
                continue
            if field_spec.field_type in CHOICE_FIELD_TYPES:
                if is_choice_field_submittable(field_result, field_spec):
                    continue
                review_reasons = set(field_result.review_reasons or [])
                reason = (
                    UNSUBMITTABLE_REVIEW_REASON
                    if UNSUBMITTABLE_REVIEW_REASON in review_reasons
                    else "choice_code_invalid"
                )
                issues.append(
                    UnsubmittableFieldIssue(
                        result_index=result_index,
                        queue_label=result.queue_label,
                        field_name=field_result.field_name,
                        field_label=field_spec.field_label,
                        value=field_result.final_value,
                        reason=reason,
                    )
                )
                continue
            if is_redcap_date_field(field_spec) and normalize_redcap_date_value(field_result.final_value) is None:
                issues.append(
                    UnsubmittableFieldIssue(
                        result_index=result_index,
                        queue_label=result.queue_label,
                        field_name=field_result.field_name,
                        field_label=field_spec.field_label,
                        value=field_result.final_value,
                        reason="invalid_date_format",
                    )
                )
                continue
            if is_redcap_numeric_field(field_spec):
                _, number_warning = normalize_redcap_numeric_value(field_result.final_value, field_spec)
                if number_warning:
                    issues.append(
                        UnsubmittableFieldIssue(
                            result_index=result_index,
                            queue_label=result.queue_label,
                            field_name=field_result.field_name,
                            field_label=field_spec.field_label,
                            value=field_result.final_value,
                            reason=number_warning.removeprefix("skipped_"),
                        )
                    )
    return issues


def is_blank_submission_value(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def is_choice_field_submittable(field_result: ExtractionFieldResult, field_spec: FieldSpec) -> bool:
    if field_spec.field_type in {"radio", "dropdown", "yesno"}:
        return single_choice_code_for_submission(field_result, field_spec) is not None
    if field_spec.field_type == "checkbox":
        if field_result.needs_review and UNSUBMITTABLE_REVIEW_REASON in set(field_result.review_reasons or []):
            return False
        return bool(checkbox_codes_for_submission(field_result, field_spec))
    return True


def normalize_checkbox_codes(result: ExtractionFieldResult) -> list[str]:
    code_value = result.final_value_code
    if isinstance(code_value, list):
        return [str(item).strip() for item in code_value if str(item).strip()]
    if isinstance(code_value, str):
        stripped = code_value.strip()
        if not stripped:
            return []
        if stripped.startswith("["):
            try:
                parsed = json.loads(stripped)
                if isinstance(parsed, list):
                    return [str(item).strip() for item in parsed if str(item).strip()]
            except json.JSONDecodeError:
                pass
        return [part.strip() for part in stripped.split(",") if part.strip()]
    value = result.final_value
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return []


def extract_tc_identity_no(result: PatientExtractionResult) -> str | None:
    for field_result in result.merged_response.results:
        if field_result.field_name != "tc_kimlik_no":
            continue
        normalized = normalize_tc_identity_no(field_result.final_value)
        if normalized:
            return normalized
    return None


def extract_record_id(result: PatientExtractionResult) -> str | None:
    for field_result in result.merged_response.results:
        if field_result.field_name != "record_id":
            continue
        if field_result.final_value in {None, ""}:
            continue
        return str(field_result.final_value).strip() or None
    return None


def normalize_tc_identity_no(value: Any) -> str | None:
    if value in {None, ""}:
        return None
    digits = "".join(ch for ch in str(value) if ch.isdigit())
    if len(digits) != 11:
        return None
    return digits
