from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, time
import json
import math
import re
import unicodedata
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError

from dictionary_parser import FieldSpec
from llm_provider import API_KEY_PROVIDER_NAMES, can_resolve_api_key, create_provider
from run_extraction import ExtractionFieldResult, ExtractionResponse, normalize_json_response_text, post_process_result
from workspace_extraction import PatientExtractionResult, merge_document_responses


NAME_VALUE_TRANSFORMS = {"name_given", "name_family", "name_given_first_2", "name_family_first_2"}


RECORD_ID_COLUMN_CANDIDATES = [
    "record_id",
    "redcap_record_id",
    "kayıt_no",
    "kayit_no",
    "kayıt numarası",
    "kayit numarasi",
]

TC_COLUMN_CANDIDATES = [
    "tc_kimlik_no",
    "t.c. kimlik no",
    "t.c. kimlik numarası",
    "tc kimlik no",
    "tc kimlik numarası",
    "tc_identity_no",
    "tc_identity",
    "tc_no",
    "tckn",
    "kimlik",
]

IDENTIFIER_COLUMN_CANDIDATES = [
    *RECORD_ID_COLUMN_CANDIDATES,
    *TC_COLUMN_CANDIDATES,
    "hasta_no",
    "patient_identifier",
    "patient_id",
    "queue_label",
]

LABEL_COLUMN_CANDIDATES = ["queue_label", "patient_label", "hasta_etiketi"]

NULL_TEXT_VALUES = {"", "nan", "none", "null", "na", "n/a"}
LLM_VALUE_NORMALIZATION_REASONS = {
    "choice_code_unmapped",
    "date_validation_failed",
    "integer_validation_failed",
    "number_validation_failed",
}
LLM_VALUE_NORMALIZATION_BATCH_SIZE = 25

GENERIC_FIELD_TOKENS = {
    "alan",
    "adet",
    "bilgi",
    "bilgileri",
    "deger",
    "degeri",
    "diger",
    "hasta",
    "hastasi",
    "no",
    "number",
    "numara",
    "numarasi",
    "sayi",
    "sayisi",
    "secim",
    "secimi",
    "sonuc",
    "sonucu",
    "veri",
}

SYNONYM_GROUPS = [
    {"ad", "adi", "isim", "name"},
    {"soyad", "soyadi", "surname", "lastname", "last"},
    {"ad", "adi", "soyad", "soyadi", "isim", "adsoyad", "adsoyadi", "isimsoyisim", "tamad", "fullname", "full"},
    {"agirlik", "kilo", "kg"},
    {"boy", "uzunluk", "cm"},
    {"dogum", "dt", "birth"},
    {"tarih", "date"},
    {"telefon", "tel", "phone"},
    {"kimlik", "tc", "tckn"},
    {"sehir", "il", "yasadigi"},
    {"biyopsi", "bx"},
    {"cocuk", "cocugu", "cocuklar", "child", "children"},
    {"ln", "lenf", "lymph", "node", "nod", "nodu"},
    {"cikarilan", "cikartilan", "diseke", "removed", "excised"},
    {"manyetik", "mr", "mri"},
    {"prostat", "psa"},
    {"patoloji", "patolojik"},
    {"ameliyat", "operasyon", "operatif", "cerrahi"},
    {"tedavi", "therapy"},
]

SYNONYM_LOOKUP: dict[str, set[str]] = {}
for group in SYNONYM_GROUPS:
    for token in group:
        SYNONYM_LOOKUP.setdefault(token, set()).update(group)


@dataclass
class ExcelColumnMapping:
    column_name: str
    field_name: str
    match_type: str
    value_transform: str = "direct"


@dataclass
class ExcelImportReport:
    results: list[PatientExtractionResult]
    mapped_columns: list[ExcelColumnMapping] = field(default_factory=list)
    ignored_columns: list[str] = field(default_factory=list)
    mapping_warnings: list[str] = field(default_factory=list)
    value_normalization_warnings: list[str] = field(default_factory=list)
    normalized_value_count: int = 0
    patient_id_column: str | None = None
    sheet_name: str | None = None


@dataclass
class ExcelImportOptions:
    sheet_name: str | None = None
    patient_id_column: str | None = None
    patient_mode: str = "auto"
    identifier_type: str | None = None
    use_llm_mapping: bool = False
    llm_settings: dict[str, Any] | None = None
    column_mappings: list[ExcelColumnMapping] = field(default_factory=list)
    automatic_mapping: bool = True
    use_llm_value_normalization: bool = False


@dataclass
class ExcelTable:
    headers: list[str]
    rows: list[dict[str, Any]]
    sheet_name: str


@dataclass
class PatientIdentity:
    patient_mode: str
    identifier_type: str | None
    identifier_value: str | None
    tc_identity_no: str | None = None


@dataclass
class PatientGroupKey:
    key: str
    value: str
    source_column: str


class LLMColumnMappingItem(BaseModel):
    column_name: str
    field_name: str
    confidence: float = 0.0
    reason: str | None = None
    value_transform: str = "direct"


class LLMColumnMappingResponse(BaseModel):
    mappings: list[LLMColumnMappingItem]


class LLMValueNormalizationItem(BaseModel):
    item_id: str
    normalized_value: Any = None
    value_code: Any = None
    confidence: float = 0.0
    reason: str | None = None


class LLMValueNormalizationResponse(BaseModel):
    items: list[LLMValueNormalizationItem]


@dataclass
class ExcelValueNormalizationTask:
    item_id: str
    field_name: str
    raw_value: Any
    review_reasons: list[str]


ProgressCallback = Callable[[int, int, str], None]


def read_excel_table(path: str | Path, sheet_name: str | None = None) -> ExcelTable:
    workbook_path = Path(path).expanduser().resolve()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        raw_rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not raw_rows:
        raise ValueError("Excel sheet is empty.")

    header_index = first_non_empty_row_index(raw_rows)
    if header_index is None:
        raise ValueError("Excel sheet does not contain a header row.")

    headers = normalize_headers(raw_rows[header_index])
    if not any(headers):
        raise ValueError("Excel header row is empty.")

    rows: list[dict[str, Any]] = []
    for raw_row in raw_rows[header_index + 1 :]:
        row = {
            header: normalize_cell_value(raw_row[index] if index < len(raw_row) else None)
            for index, header in enumerate(headers)
            if header
        }
        if any(value is not None for value in row.values()):
            rows.append(row)

    if not rows:
        raise ValueError("Excel sheet does not contain data rows.")

    return ExcelTable(headers=[header for header in headers if header], rows=rows, sheet_name=str(worksheet.title))


def import_patient_excel(
    *,
    path: str | Path,
    project_name: str,
    field_specs_by_name: dict[str, FieldSpec],
    options: ExcelImportOptions | None = None,
    progress_callback: ProgressCallback | None = None,
) -> ExcelImportReport:
    import_options = options or ExcelImportOptions()
    if progress_callback is not None:
        progress_callback(0, 0, "mapping")
    table = read_excel_table(path, import_options.sheet_name)
    patient_id_column = resolve_patient_id_column(
        headers=table.headers,
        requested_column=import_options.patient_id_column,
    )
    label_column = resolve_optional_column(table.headers, LABEL_COLUMN_CANDIDATES)
    excluded_mapping_columns = {patient_id_column}
    if label_column:
        excluded_mapping_columns.add(label_column)
    column_mappings = sanitize_column_mappings(
        import_options.column_mappings,
        headers=table.headers,
        field_specs_by_name=field_specs_by_name,
        excluded_columns=excluded_mapping_columns,
    )
    if import_options.automatic_mapping:
        column_mappings = merge_column_mappings(
            column_mappings,
            build_procedural_column_mappings(
                table.headers,
                table.rows,
                field_specs_by_name,
                excluded_columns=excluded_mapping_columns,
            ),
        )
    mapping_warnings: list[str] = []
    unmapped_headers = unresolved_mapping_headers(
        table.headers,
        column_mappings,
        excluded_columns=excluded_mapping_columns,
    )
    if import_options.use_llm_mapping and unmapped_headers:
        if not import_options.llm_settings:
            mapping_warnings.append("llm_mapping_skipped_missing_settings")
        elif not can_call_llm(import_options.llm_settings):
            mapping_warnings.append("llm_mapping_skipped_missing_api_key")
        else:
            try:
                llm_mappings = build_llm_column_mappings(
                    headers=unmapped_headers,
                    rows=table.rows,
                    field_specs_by_name=field_specs_by_name,
                    existing_mappings=column_mappings,
                    llm_settings=import_options.llm_settings,
                )
            except ValueError as exc:
                llm_mappings = []
                mapping_warnings.append(f"llm_mapping_failed:{str(exc)[:300]}")
            column_mappings = merge_column_mappings(column_mappings, llm_mappings)
    mapped_column_names = {mapping.column_name for mapping in column_mappings}
    requested_identifier_type = import_options.identifier_type

    grouped_rows: dict[str, list[dict[str, Any]]] = {}
    group_keys: dict[str, PatientGroupKey] = {}
    for row in table.rows:
        group_key = build_row_group_key(
            row=row,
            primary_column=patient_id_column,
            headers=table.headers,
            column_mappings=column_mappings,
        )
        if group_key is None:
            continue
        group_keys[group_key.key] = group_key
        grouped_rows.setdefault(group_key.key, []).append(row)

    if not grouped_rows:
        raise ValueError(f"Excel sheet does not contain patient identifiers in column '{patient_id_column}'.")

    grouped_row_responses: dict[str, list[ExtractionResponse]] = {}
    for group_key, rows in grouped_rows.items():
        row_responses: list[ExtractionResponse] = []
        for row_index, row in enumerate(rows, start=1):
            row_responses.append(
                build_row_response(
                    project_name=project_name,
                    row=row,
                    row_index=row_index,
                    mappings=column_mappings,
                    field_specs_by_name=field_specs_by_name,
                )
            )
        grouped_row_responses[group_key] = row_responses

    value_normalization_warnings: list[str] = []
    normalized_value_count = 0
    if import_options.use_llm_value_normalization:
        if not import_options.llm_settings:
            value_normalization_warnings.append("llm_value_normalization_skipped_missing_settings")
        elif not can_call_llm(import_options.llm_settings):
            value_normalization_warnings.append("llm_value_normalization_skipped_missing_api_key")
        else:
            all_row_responses = [
                response
                for row_responses in grouped_row_responses.values()
                for response in row_responses
            ]
            try:
                normalized_value_count = normalize_excel_values_with_llm(
                    responses=all_row_responses,
                    field_specs_by_name=field_specs_by_name,
                    llm_settings=import_options.llm_settings,
                    progress_callback=progress_callback,
                )
            except ValueError as exc:
                value_normalization_warnings.append(f"llm_value_normalization_failed:{str(exc)[:300]}")

    if progress_callback is not None:
        progress_callback(0, 0, "building_patients")

    results: list[PatientExtractionResult] = []
    for group_key, rows in grouped_rows.items():
        patient_group = group_keys[group_key]
        row_responses = grouped_row_responses.get(group_key, [])
        merged_response = merge_document_responses(project_name=project_name, responses=row_responses)
        patient_identity = resolve_patient_identity(
            rows=rows,
            patient_key=patient_group.value,
            patient_id_column=patient_group.source_column,
            requested_patient_mode=import_options.patient_mode,
            requested_identifier_type=requested_identifier_type,
            column_mappings=column_mappings,
        )
        if patient_identity.tc_identity_no:
            ensure_tc_identifier_result(
                response=merged_response,
                tc_identity_no=patient_identity.tc_identity_no,
                field_specs_by_name=field_specs_by_name,
            )
        queue_label = build_queue_label(
            rows=rows,
            label_column=label_column,
            patient_id_column=patient_group.source_column,
            patient_key=patient_group.value,
        )
        results.append(
            PatientExtractionResult(
                queue_label=queue_label,
                patient_mode=patient_identity.patient_mode,
                identifier_type=patient_identity.identifier_type,
                identifier_value=patient_identity.identifier_value,
                documents=[str(Path(path).expanduser().resolve())],
                document_results=[
                    (f"{Path(path).name}:{table.sheet_name}:row-{index}", response)
                    for index, response in enumerate(row_responses, start=1)
                ],
                merged_response=merged_response,
            )
        )

    ignored_columns = [
        header
        for header in table.headers
        if header not in mapped_column_names and header not in {patient_id_column, label_column}
    ]
    return ExcelImportReport(
        results=results,
        mapped_columns=column_mappings,
        ignored_columns=ignored_columns,
        mapping_warnings=mapping_warnings,
        value_normalization_warnings=value_normalization_warnings,
        normalized_value_count=normalized_value_count,
        patient_id_column=patient_id_column,
        sheet_name=table.sheet_name,
    )


def build_procedural_column_mappings(
    headers: list[str],
    rows: list[dict[str, Any]],
    field_specs_by_name: dict[str, FieldSpec],
    excluded_columns: set[str] | None = None,
    allow_composite_name_splits: bool = False,
) -> list[ExcelColumnMapping]:
    column_mappings = build_excel_column_mappings(
        headers,
        field_specs_by_name,
        excluded_columns=excluded_columns,
    )
    column_mappings = merge_column_mappings(
        column_mappings,
        build_fast_semantic_column_mappings(
            headers,
            field_specs_by_name,
            existing_mappings=column_mappings,
            excluded_columns=excluded_columns,
        ),
    )
    if not allow_composite_name_splits:
        return column_mappings
    return merge_column_mappings(
        column_mappings,
        build_fast_person_name_column_mappings(
            headers,
            rows,
            field_specs_by_name,
            existing_mappings=column_mappings,
            excluded_columns=excluded_columns,
        ),
    )


def sanitize_column_mappings(
    mappings: list[ExcelColumnMapping],
    *,
    headers: list[str],
    field_specs_by_name: dict[str, FieldSpec],
    excluded_columns: set[str] | None = None,
) -> list[ExcelColumnMapping]:
    excluded = excluded_columns or set()
    header_set = set(headers)
    field_set = set(field_specs_by_name)
    sanitized: list[ExcelColumnMapping] = []
    used_pairs: set[tuple[str, str]] = set()
    used_fields: set[str] = set()
    for mapping in mappings:
        if mapping.column_name not in header_set or mapping.column_name in excluded:
            continue
        if mapping.field_name not in field_set:
            continue
        pair = (mapping.column_name, mapping.field_name)
        if pair in used_pairs or mapping.field_name in used_fields:
            continue
        used_pairs.add(pair)
        used_fields.add(mapping.field_name)
        sanitized.append(
            ExcelColumnMapping(
                column_name=mapping.column_name,
                field_name=mapping.field_name,
                match_type=mapping.match_type or "manual",
                value_transform=normalize_value_transform(mapping.value_transform),
            )
        )
    return sanitized


def build_excel_column_mappings(
    headers: list[str],
    field_specs_by_name: dict[str, FieldSpec],
    excluded_columns: set[str] | None = None,
) -> list[ExcelColumnMapping]:
    excluded = excluded_columns or set()
    exact_field_names = {field_name: field_name for field_name in field_specs_by_name}
    normalized_field_names = unique_lookup(
        (normalize_match_key(field_name), field_name) for field_name in field_specs_by_name
    )
    normalized_labels = unique_lookup(
        (normalize_match_key(spec.field_label), spec.field_name)
        for spec in field_specs_by_name.values()
        if spec.field_label
    )

    mappings: list[ExcelColumnMapping] = []
    used_fields: set[str] = set()
    for header in headers:
        if header in excluded:
            continue
        field_name = exact_field_names.get(header)
        match_type = "field_name"
        if field_name is None:
            field_name = normalized_field_names.get(normalize_match_key(header))
            match_type = "normalized_field_name"
        if field_name is None:
            field_name = normalized_labels.get(normalize_match_key(header))
            match_type = "field_label"
        if field_name is None:
            field_name = find_embedded_field_name(header, normalized_field_names)
            match_type = "embedded_field_name"
        if field_name is None or field_name in used_fields:
            continue
        used_fields.add(field_name)
        mappings.append(ExcelColumnMapping(column_name=header, field_name=field_name, match_type=match_type))
    return mappings


def build_fast_semantic_column_mappings(
    headers: list[str],
    field_specs_by_name: dict[str, FieldSpec],
    *,
    existing_mappings: list[ExcelColumnMapping],
    excluded_columns: set[str] | None = None,
) -> list[ExcelColumnMapping]:
    excluded = excluded_columns or set()
    used_columns = {mapping.column_name for mapping in existing_mappings}
    used_fields = {mapping.field_name for mapping in existing_mappings}
    mappings: list[ExcelColumnMapping] = []
    for header in headers:
        if header in excluded or header in used_columns:
            continue
        ranked = ranked_field_candidates_for_header(
            header=header,
            field_specs_by_name={
                field_name: spec
                for field_name, spec in field_specs_by_name.items()
                if field_name not in used_fields
            },
            limit=2,
        )
        if not ranked:
            continue
        if should_defer_composite_name_header_to_llm(header, ranked, field_specs_by_name):
            continue
        best_field, best_score = ranked[0]
        next_score = ranked[1][1] if len(ranked) > 1 else 0.0
        if best_score >= 0.78 and best_score - next_score >= 0.08:
            if not is_column_mapping_acceptable(
                column_name=header,
                field_spec=field_specs_by_name[best_field],
                rows=[],
                value_transform="direct",
            ):
                continue
            mappings.append(
                ExcelColumnMapping(
                    column_name=header,
                    field_name=best_field,
                    match_type="semantic_header",
                )
            )
            used_columns.add(header)
            used_fields.add(best_field)
    return mappings


def build_fast_person_name_column_mappings(
    headers: list[str],
    rows: list[dict[str, Any]],
    field_specs_by_name: dict[str, FieldSpec],
    *,
    existing_mappings: list[ExcelColumnMapping],
    excluded_columns: set[str] | None = None,
) -> list[ExcelColumnMapping]:
    excluded = excluded_columns or set()
    used_columns = {mapping.column_name for mapping in existing_mappings}
    used_fields = {mapping.field_name for mapping in existing_mappings}
    mappings: list[ExcelColumnMapping] = []
    for header in headers:
        if header in excluded or header in used_columns:
            continue
        if not column_looks_like_generic_name_column(header):
            continue
        if not sample_values_look_like_person_names(rows, header):
            continue

        given_field = best_name_field(
            header=header,
            field_specs_by_name=field_specs_by_name,
            used_fields=used_fields,
            predicate=field_looks_like_given_name,
        )
        family_field = best_name_field(
            header=header,
            field_specs_by_name=field_specs_by_name,
            used_fields=used_fields,
            predicate=field_looks_like_family_name,
        )
        if given_field:
            mappings.append(
                ExcelColumnMapping(
                    column_name=header,
                    field_name=given_field,
                    match_type="semantic_full_name",
                    value_transform="name_given",
                )
            )
            used_fields.add(given_field)
        if family_field and sample_values_look_like_composite_person_names(rows, header):
            mappings.append(
                ExcelColumnMapping(
                    column_name=header,
                    field_name=family_field,
                    match_type="semantic_full_name",
                    value_transform="name_family",
                )
            )
            used_fields.add(family_field)
    return mappings


def best_name_field(
    *,
    header: str,
    field_specs_by_name: dict[str, FieldSpec],
    used_fields: set[str],
    predicate: Any,
) -> str | None:
    candidates = [
        (field_name, score_header_to_field(header, spec))
        for field_name, spec in field_specs_by_name.items()
        if field_name not in used_fields
        and spec.field_type in {"text", "notes"}
        and predicate(spec)
    ]
    if not candidates:
        return None
    return sorted(candidates, key=lambda item: item[1], reverse=True)[0][0]


def should_defer_composite_name_header_to_llm(
    header: str,
    ranked: list[tuple[str, float]],
    field_specs_by_name: dict[str, FieldSpec],
) -> bool:
    header_key = normalize_match_key(header)
    if header_key not in {"isim", "adsoyad", "adsoyadi", "isimsoyisim", "tamad", "fullname", "full", "name"}:
        return False
    candidates = [field_specs_by_name[field_name] for field_name, score in ranked if score >= 0.4]
    has_given = any(field_looks_like_given_name(spec) for spec in candidates)
    has_family = any(field_looks_like_family_name(spec) for spec in candidates)
    return has_given and has_family


def field_looks_like_given_name(field_spec: FieldSpec) -> bool:
    text = normalize_search_text(" ".join([field_spec.field_name, clean_metadata_text(field_spec.field_label)]))
    tokens = set(text.split())
    joined = "".join(tokens)
    has_middle_marker = bool(tokens & {"gobek", "middle", "ikinci"}) or "gobekad" in joined or "middlename" in joined
    return (
        bool(tokens & {"ad", "adi", "isim", "name"})
        and not bool(tokens & {"soyad", "soyadi", "surname"})
        and not has_middle_marker
    )


def field_looks_like_family_name(field_spec: FieldSpec) -> bool:
    text = normalize_search_text(" ".join([field_spec.field_name, clean_metadata_text(field_spec.field_label)]))
    tokens = set(text.split())
    return bool(tokens & {"soyad", "soyadi", "surname", "lastname", "last"})


def field_looks_like_middle_name(field_spec: FieldSpec) -> bool:
    text = normalize_search_text(" ".join([field_spec.field_name, clean_metadata_text(field_spec.field_label)]))
    tokens = set(text.split())
    joined = "".join(tokens)
    return bool(tokens & {"gobek", "middle", "ikinci"}) or "gobekad" in joined or "middlename" in joined


def field_looks_like_person_name(field_spec: FieldSpec) -> bool:
    return (
        field_looks_like_given_name(field_spec)
        or field_looks_like_family_name(field_spec)
        or field_looks_like_middle_name(field_spec)
    )


def build_llm_column_mappings(
    *,
    headers: list[str],
    rows: list[dict[str, Any]],
    field_specs_by_name: dict[str, FieldSpec],
    existing_mappings: list[ExcelColumnMapping],
    llm_settings: dict[str, Any],
) -> list[ExcelColumnMapping]:
    already_mapped_columns = {mapping.column_name for mapping in existing_mappings}
    already_mapped_fields = {mapping.field_name for mapping in existing_mappings}
    columns_payload = [
        {
            "column_name": header,
            "candidate_fields": candidate_field_payload_for_header(
                header=header,
                rows=rows,
                field_specs_by_name={
                    field_name: spec
                    for field_name, spec in field_specs_by_name.items()
                    if field_name not in already_mapped_fields
                },
            ),
            "sample_values_if_needed": sample_column_values(rows, header, limit=3),
            "already_mapped": header in already_mapped_columns,
        }
        for header in headers
    ]
    messages = [
        {
            "role": "system",
            "content": (
                "You map custom clinical Excel columns to REDCap data dictionary fields for a clinical import workflow. "
                "Return JSON only. Use column names first and sample values only when the name is ambiguous. "
                "Choose only from candidate_fields. Do not invent field names. "
                "One Excel column may map to multiple REDCap fields when the column contains a composite value, "
                "for example a Turkish 'İsim'/'Ad Soyad' column can map to both first-name and surname fields."
            ),
        },
        {
            "role": "user",
            "content": json.dumps(
                {
                    "task": "Map Excel columns to REDCap fields.",
                    "rules": [
                        "Use the Excel column name as the primary evidence.",
                        "Use sample_values_if_needed only if the name is ambiguous.",
                        "A custom Turkish clinic heading may be a synonym or shorthand for a REDCap field.",
                        "candidate_fields may include low score_hint fields so that you can still recognize custom clinical synonyms.",
                        "If a column contains a full name and REDCap has separate given-name and surname fields, return two mappings for the same column.",
                        "Use value_transform='name_given' for the given-name/ad field and value_transform='name_family' for the surname/soyad field.",
                        "When a REDCap field stores only the first two letters, value_transform may be 'first_2', 'name_given_first_2', or 'name_family_first_2'.",
                        "A generic column named 'İsim', 'Ad Soyad', 'Name', or 'Full name' must not be mapped to a middle-name/göbek-ad field unless the column name explicitly says göbek, middle, or ikinci ad.",
                        "Never map procedure, operation, surgery, treatment, biopsy, pathology, stage, result, or follow-up columns to patient name/identity fields.",
                        "Never map lymph node/LN/nodal count columns to demographic child/children count fields; 'LN sayısı' means lymph node count, not child count.",
                        "A confident mapping requires that the column name and sample values are compatible with the REDCap field meaning.",
                        "Use value_transform='direct' for ordinary one-to-one mappings.",
                        "Return only confident mappings. Leave unrelated columns unmapped.",
                        "Do not map a column to a field outside candidate_fields.",
                    ],
                    "allowed_value_transform": [
                        "direct",
                        "name_given",
                        "name_family",
                        "first_2",
                        "name_given_first_2",
                        "name_family_first_2",
                    ],
                    "excel_columns": columns_payload,
                    "response_schema": {
                        "mappings": [
                            {
                                "column_name": "exact Excel column name",
                                "field_name": "exact REDCap field_name",
                                "confidence": 0.0,
                                "value_transform": "direct | name_given | name_family | first_2 | name_given_first_2 | name_family_first_2",
                                "reason": "short reason",
                            }
                        ]
                    },
                },
                ensure_ascii=False,
            ),
        },
    ]
    settings = dict(llm_settings)
    provider = create_provider(settings)
    schema = LLMColumnMappingResponse.model_json_schema()
    response = provider.generate(messages, schema, settings)
    content = response.content
    if not content and settings.get("think", False):
        retry_settings = dict(settings)
        retry_settings["think"] = False
        content = create_provider(retry_settings).generate(messages, schema, retry_settings).content
    if not content:
        raise ValueError("LLM column mapping returned an empty response.")

    parsed = parse_llm_column_mapping_response(
        content=content,
        headers=headers,
        field_specs_by_name=field_specs_by_name,
    )

    header_set = set(headers)
    field_set = set(field_specs_by_name)
    mappings: list[ExcelColumnMapping] = []
    for item in parsed.mappings:
        if item.column_name not in header_set:
            continue
        if item.field_name not in field_set:
            continue
        if float(item.confidence or 0.0) < 0.55:
            continue
        value_transform = normalize_value_transform(item.value_transform)
        if not is_column_mapping_acceptable(
            column_name=item.column_name,
            field_spec=field_specs_by_name[item.field_name],
            rows=rows,
            value_transform=value_transform,
        ):
            continue
        mappings.append(
            ExcelColumnMapping(
                column_name=item.column_name,
                field_name=item.field_name,
                match_type="llm",
                value_transform=value_transform,
            )
        )
    return mappings


def candidate_field_payload_for_header(
    *,
    header: str,
    rows: list[dict[str, Any]] | None = None,
    field_specs_by_name: dict[str, FieldSpec],
    limit: int = 24,
) -> list[dict[str, Any]]:
    sample_rows = rows or []
    header_concepts = semantic_concepts_for_text(header)
    if len(field_specs_by_name) <= 60:
        ranked = sorted(
            [
                (field_name, score_header_to_field(header, spec))
                for field_name, spec in field_specs_by_name.items()
            ],
            key=lambda item: item[1],
            reverse=True,
        )
    else:
        ranked = ranked_field_candidates_for_header(
            header=header,
            field_specs_by_name=field_specs_by_name,
            limit=limit,
        )
    return [
        {
            "field_name": clean_metadata_text(field_name),
            "form_name": clean_metadata_text(field_specs_by_name[field_name].form_name),
            "field_type": clean_metadata_text(field_specs_by_name[field_name].field_type),
            "field_label": clean_metadata_text(field_specs_by_name[field_name].field_label),
            "field_note": clean_metadata_text(field_specs_by_name[field_name].field_note),
            "text_validation": clean_metadata_text(field_specs_by_name[field_name].text_validation),
            "score_hint": round(score, 3),
        }
        for field_name, score in ranked
        if should_include_llm_candidate_field(
            header=header,
            header_concepts=header_concepts,
            rows=sample_rows,
            field_spec=field_specs_by_name[field_name],
            score=score,
            allow_low_score=len(field_specs_by_name) <= 60,
        )
    ]


def should_include_llm_candidate_field(
    *,
    header: str,
    header_concepts: set[str],
    rows: list[dict[str, Any]],
    field_spec: FieldSpec,
    score: float,
    allow_low_score: bool,
) -> bool:
    if score > 0:
        return True
    if not allow_low_score:
        return False
    if not is_column_mapping_acceptable(
        column_name=header,
        field_spec=field_spec,
        rows=rows,
        value_transform="direct",
    ):
        return False
    if header_concepts:
        return bool(
            header_concepts
            & semantic_concepts_for_text(field_semantic_text(field_spec))
        )
    return True


def ranked_field_candidates_for_header(
    *,
    header: str,
    field_specs_by_name: dict[str, FieldSpec],
    limit: int,
) -> list[tuple[str, float]]:
    scored = [
        (field_name, score_header_to_field(header, spec))
        for field_name, spec in field_specs_by_name.items()
    ]
    scored = [(field_name, score) for field_name, score in scored if score > 0]
    return sorted(scored, key=lambda item: item[1], reverse=True)[:limit]


def score_header_to_field(header: str, field_spec: FieldSpec) -> float:
    header_key = normalize_match_key(header)
    if not is_column_mapping_acceptable(
        column_name=header,
        field_spec=field_spec,
        rows=[],
        value_transform="direct",
    ):
        return 0.0
    field_name_text = clean_metadata_text(field_spec.field_name)
    field_label_text = clean_metadata_text(field_spec.field_label)
    field_note_text = clean_metadata_text(field_spec.field_note)
    text_validation_text = clean_metadata_text(field_spec.text_validation)
    field_name_key = normalize_match_key(field_name_text)
    label_key = normalize_match_key(field_label_text)
    note_key = normalize_match_key(field_note_text)
    if not header_key:
        return 0.0
    if header_key in {field_name_key, label_key}:
        return 1.0
    if len(header_key) >= 4 and (header_key in field_name_key or field_name_key in header_key):
        return 0.92
    if len(header_key) >= 4 and (header_key in label_key or label_key in header_key):
        return 0.90

    header_tokens = tokenize_match_text(header)
    field_tokens = tokenize_match_text(
        " ".join(
            [
                field_name_text,
                field_label_text,
                field_note_text,
                text_validation_text,
            ]
        )
    )
    if not header_tokens or not field_tokens:
        return 0.0

    overlap = header_tokens & field_tokens
    precision = len(overlap) / len(header_tokens)
    recall = len(overlap) / len(field_tokens)
    score = (precision * 0.65) + (recall * 0.25)

    if acronym_for_text(field_label_text) == header_key and len(header_key) >= 2:
        score = max(score, 0.86)
    if header_key and header_key in note_key and len(header_key) >= 4:
        score = max(score, 0.72)
    return min(score, 1.0)


def parse_llm_column_mapping_response(
    *,
    content: str,
    headers: list[str],
    field_specs_by_name: dict[str, FieldSpec],
) -> LLMColumnMappingResponse:
    normalized = normalize_json_response_text(content)
    try:
        return LLMColumnMappingResponse.model_validate_json(normalized)
    except ValidationError as json_exc:
        json_fallback = parse_llm_mapping_json_fallback(normalized)
        if json_fallback is not None:
            return json_fallback
        text_fallback = parse_llm_mapping_text_fallback(
            content=content,
            headers=headers,
            field_specs_by_name=field_specs_by_name,
        )
        if text_fallback.mappings:
            return text_fallback
        raise ValueError(f"invalid_json:{json_exc}") from json_exc


def parse_llm_mapping_json_fallback(content: str) -> LLMColumnMappingResponse | None:
    candidates = [content.strip()]
    first_object = content.find("{")
    last_object = content.rfind("}")
    if first_object != -1 and last_object > first_object:
        candidates.append(content[first_object : last_object + 1].strip())
    first_array = content.find("[")
    last_array = content.rfind("]")
    if first_array != -1 and last_array > first_array:
        candidates.append(content[first_array : last_array + 1].strip())

    for candidate in candidates:
        try:
            raw = json.loads(candidate)
        except json.JSONDecodeError:
            continue
        if isinstance(raw, dict) and isinstance(raw.get("mappings"), list):
            return LLMColumnMappingResponse.model_validate(raw)
        if isinstance(raw, list):
            return LLMColumnMappingResponse.model_validate({"mappings": raw})
    return None


def parse_llm_mapping_text_fallback(
    *,
    content: str,
    headers: list[str],
    field_specs_by_name: dict[str, FieldSpec],
) -> LLMColumnMappingResponse:
    header_by_key = {normalize_match_key(header): header for header in headers}
    field_names = set(field_specs_by_name)
    mappings: list[LLMColumnMappingItem] = []

    for line in content.splitlines():
        line_text = line.strip().strip("-* ")
        if not line_text:
            continue
        field_name = extract_field_name_from_text(line_text, field_names)
        if not field_name:
            continue
        column_name = extract_column_name_from_text(line_text, headers, header_by_key, field_name)
        if not column_name:
            continue
        mappings.append(
            LLMColumnMappingItem(
                column_name=column_name,
                field_name=field_name,
                confidence=0.70,
                reason="parsed_from_non_json_llm_output",
            )
        )

    return LLMColumnMappingResponse(mappings=dedupe_llm_mapping_items(mappings))


def extract_field_name_from_text(line: str, field_names: set[str]) -> str | None:
    bracket_values = re.findall(r"\[([A-Za-z0-9_]+)\]", line)
    for value in bracket_values:
        if value in field_names:
            return value

    normalized_line = normalize_match_key(line)
    matches = [field_name for field_name in field_names if normalize_match_key(field_name) in normalized_line]
    if len(matches) == 1:
        return matches[0]
    if matches:
        return sorted(matches, key=len, reverse=True)[0]
    return None


def extract_column_name_from_text(
    line: str,
    headers: list[str],
    header_by_key: dict[str, str],
    field_name: str,
) -> str | None:
    for header in headers:
        if header and header in line:
            return header

    normalized_line = normalize_match_key(line)
    header_matches = [
        header
        for key, header in header_by_key.items()
        if key and key in normalized_line and normalize_match_key(header) != normalize_match_key(field_name)
    ]
    if len(header_matches) == 1:
        return header_matches[0]
    if header_matches:
        return sorted(header_matches, key=len, reverse=True)[0]

    stripped = re.sub(r"\[[A-Za-z0-9_]+\]", " ", line)
    stripped = re.sub(r"\b(to|field|maps?|mapped|redcap|alan|eşleşir|eslesir)\b", " ", stripped, flags=re.IGNORECASE)
    candidate_key = normalize_match_key(stripped)
    return header_by_key.get(candidate_key)


def dedupe_llm_mapping_items(items: list[LLMColumnMappingItem]) -> list[LLMColumnMappingItem]:
    deduped: list[LLMColumnMappingItem] = []
    used_fields: set[str] = set()
    used_pairs: set[tuple[str, str]] = set()
    for item in items:
        pair = (item.column_name, item.field_name)
        if pair in used_pairs or item.field_name in used_fields:
            continue
        used_pairs.add(pair)
        used_fields.add(item.field_name)
        item.value_transform = normalize_value_transform(item.value_transform)
        deduped.append(item)
    return deduped


def normalize_value_transform(value: str | None) -> str:
    normalized = str(value or "direct").strip().lower()
    aliases = {
        "given_name": "name_given",
        "first_name": "name_given",
        "ad": "name_given",
        "family_name": "name_family",
        "last_name": "name_family",
        "surname": "name_family",
        "soyad": "name_family",
    }
    normalized = aliases.get(normalized, normalized)
    aliases.update(
        {
            "first2": "first_2",
            "first_two": "first_2",
            "ilk2": "first_2",
            "given_name_first_2": "name_given_first_2",
            "first_name_first_2": "name_given_first_2",
            "ad_ilk_2": "name_given_first_2",
            "family_name_first_2": "name_family_first_2",
            "last_name_first_2": "name_family_first_2",
            "surname_first_2": "name_family_first_2",
            "soyad_ilk_2": "name_family_first_2",
        }
    )
    normalized = aliases.get(normalized, normalized)
    allowed = {"direct", "name_given", "name_family", "first_2", "name_given_first_2", "name_family_first_2"}
    return normalized if normalized in allowed else "direct"


def is_column_mapping_acceptable(
    *,
    column_name: str,
    field_spec: FieldSpec,
    rows: list[dict[str, Any]],
    value_transform: str,
) -> bool:
    if not field_column_concepts_compatible(column_name, field_spec):
        return False
    if field_looks_like_middle_name(field_spec) and not column_looks_like_middle_name(column_name):
        return False
    if field_looks_like_person_name(field_spec):
        if column_has_non_person_concept(column_name):
            return False
        if not (
            column_looks_like_person_name(column_name)
            or sample_values_look_like_person_names(rows, column_name)
            or value_transform in NAME_VALUE_TRANSFORMS
        ):
            return False
    if value_transform in NAME_VALUE_TRANSFORMS:
        if field_spec.field_type not in {"text", "notes"}:
            return False
        if not (column_looks_like_person_name(column_name) or sample_values_look_like_person_names(rows, column_name)):
            return False
    return True


def is_llm_mapping_acceptable(
    *,
    column_name: str,
    field_spec: FieldSpec,
    rows: list[dict[str, Any]],
    value_transform: str,
) -> bool:
    return is_column_mapping_acceptable(
        column_name=column_name,
        field_spec=field_spec,
        rows=rows,
        value_transform=value_transform,
    )


def field_column_concepts_compatible(column_name: str, field_spec: FieldSpec) -> bool:
    column_concepts = semantic_concepts_for_text(column_name)
    field_concepts = semantic_concepts_for_text(field_semantic_text(field_spec))
    if not column_concepts or not field_concepts:
        return True

    exclusive_concepts = {"child", "lymph_node"}
    column_exclusive = column_concepts & exclusive_concepts
    field_exclusive = field_concepts & exclusive_concepts
    if column_exclusive and field_exclusive and not (column_exclusive & field_exclusive):
        return False

    clinical_concepts = {"biopsy", "lymph_node", "operation", "pathology", "prostate", "stage", "treatment"}
    if "child" in field_concepts and column_concepts & clinical_concepts and "child" not in column_concepts:
        return False
    if "lymph_node" in field_concepts and "child" in column_concepts and "lymph_node" not in column_concepts:
        return False
    return True


def field_semantic_text(field_spec: FieldSpec) -> str:
    return " ".join(
        [
            field_spec.field_name,
            clean_metadata_text(field_spec.field_label),
            clean_metadata_text(field_spec.field_note),
        ]
    )


def semantic_concepts_for_text(value: Any) -> set[str]:
    text = clean_metadata_text(value)
    tokens = set(normalize_search_text(text).split())
    key = normalize_match_key(text)
    concepts: set[str] = set()
    if tokens & {"cocuk", "cocugu", "cocuklar", "child", "children"}:
        concepts.add("child")
    if tokens & {"ln", "lenf", "lymph", "node", "nod", "nodu"} or any(
        marker in key for marker in {"lenfnodu", "lymphnode", "lymphnodes"}
    ):
        concepts.add("lymph_node")
    if tokens & {"ameliyat", "cerrahi", "operasyon", "operatif", "prostatektomi"}:
        concepts.add("operation")
    if tokens & {"biyopsi", "bx"}:
        concepts.add("biopsy")
    if tokens & {"patoloji", "patolojik", "histoloji", "histopatolojik"}:
        concepts.add("pathology")
    if tokens & {"evre", "stage", "staging"}:
        concepts.add("stage")
    if tokens & {"tedavi", "therapy"}:
        concepts.add("treatment")
    if tokens & {"prostat", "prostate", "psa"}:
        concepts.add("prostate")
    return concepts


def column_looks_like_person_name(column_name: str) -> bool:
    tokens = set(normalize_search_text(column_name).split())
    joined = "".join(tokens)
    return bool(tokens & {"ad", "adi", "isim", "soyad", "soyadi", "name", "surname", "fullname"}) or any(
        marker in joined
        for marker in {"adsoyad", "adsoyadi", "isimsoyisim", "tamad", "fullname"}
    )


def column_looks_like_generic_name_column(column_name: str) -> bool:
    key = normalize_match_key(column_name)
    if key in {"isim", "ad", "adi", "adsoyad", "adsoyadi", "isimsoyisim", "tamad", "fullname", "full", "name"}:
        return True
    tokens = set(normalize_search_text(column_name).split())
    return bool(tokens & {"isim", "name", "fullname"}) or bool(tokens & {"ad", "adi"} and tokens & {"soyad", "soyadi"})


def column_looks_like_middle_name(column_name: str) -> bool:
    tokens = set(normalize_search_text(column_name).split())
    joined = "".join(tokens)
    return bool(tokens & {"gobek", "middle", "ikinci"}) or "gobekad" in joined or "middlename" in joined


def column_has_non_person_concept(column_name: str) -> bool:
    tokens = set(normalize_search_text(column_name).split())
    return bool(
        tokens
        & {
            "ameliyat",
            "biyopsi",
            "cerrahi",
            "durum",
            "evre",
            "gleason",
            "histoloji",
            "histopatolojik",
            "izlem",
            "operasyon",
            "patoloji",
            "prostatektomi",
            "radikal",
            "sonuc",
            "tedavi",
            "tip",
            "tipi",
        }
    )


def sample_values_look_like_person_names(rows: list[dict[str, Any]], column_name: str) -> bool:
    values = [normalize_cell_value(row.get(column_name)) for row in rows[:12]]
    texts = [scalar_to_text(value).strip() for value in values if value is not None]
    if not texts:
        return False
    person_like = 0
    checked = 0
    for text in texts:
        if not text or any(char.isdigit() for char in text):
            continue
        tokens = [token for token in re.split(r"\s+", text) if token]
        if not tokens or len(tokens) > 5:
            continue
        checked += 1
        if all(re.search(r"[A-Za-zÇĞİÖŞÜçğıöşü]", token) for token in tokens):
            person_like += 1
    return checked > 0 and person_like / checked >= 0.65


def sample_values_look_like_composite_person_names(rows: list[dict[str, Any]], column_name: str) -> bool:
    values = [normalize_cell_value(row.get(column_name)) for row in rows[:12]]
    texts = [scalar_to_text(value).strip() for value in values if value is not None]
    if not texts:
        return False
    person_like = 0
    checked = 0
    for text in texts:
        if not text or any(char.isdigit() for char in text):
            continue
        tokens = [token for token in re.split(r"\s+", text) if token]
        if len(tokens) < 2 or len(tokens) > 5:
            checked += 1
            continue
        checked += 1
        if all(re.search(r"[A-Za-zÇĞİÖŞÜçğıöşü]", token) for token in tokens):
            person_like += 1
    return checked > 0 and person_like / checked >= 0.65


def merge_column_mappings(
    primary: list[ExcelColumnMapping],
    secondary: list[ExcelColumnMapping],
) -> list[ExcelColumnMapping]:
    merged: list[ExcelColumnMapping] = []
    used_fields: set[str] = set()
    used_pairs: set[tuple[str, str]] = set()
    for mapping in [*primary, *secondary]:
        pair = (mapping.column_name, mapping.field_name)
        if pair in used_pairs or mapping.field_name in used_fields:
            continue
        used_pairs.add(pair)
        used_fields.add(mapping.field_name)
        merged.append(mapping)
    return merged


def unresolved_mapping_headers(
    headers: list[str],
    mappings: list[ExcelColumnMapping],
    *,
    excluded_columns: set[str] | None = None,
) -> list[str]:
    excluded = excluded_columns or set()
    mapped = {mapping.column_name for mapping in mappings}
    return [
        header
        for header in headers
        if header not in excluded and header not in mapped
    ]


def can_call_llm(settings: dict[str, Any]) -> bool:
    provider_name = str(settings.get("provider", "ollama")).strip().lower()
    if provider_name == "ollama":
        return True
    if provider_name not in API_KEY_PROVIDER_NAMES:
        return True
    return can_resolve_api_key(settings)


def normalize_excel_values_with_llm(
    *,
    responses: list[ExtractionResponse],
    field_specs_by_name: dict[str, FieldSpec],
    llm_settings: dict[str, Any],
    progress_callback: ProgressCallback | None = None,
) -> int:
    tasks = collect_llm_value_normalization_tasks(
        responses=responses,
        field_specs_by_name=field_specs_by_name,
    )
    if progress_callback is not None:
        progress_callback(0, len(tasks), "normalizing_values")
    if not tasks:
        return 0

    normalizations: dict[tuple[str, str], LLMValueNormalizationItem] = {}
    settings = build_value_normalization_llm_settings(llm_settings)
    provider = create_provider(settings)
    schema = LLMValueNormalizationResponse.model_json_schema()

    for offset in range(0, len(tasks), LLM_VALUE_NORMALIZATION_BATCH_SIZE):
        batch = tasks[offset : offset + LLM_VALUE_NORMALIZATION_BATCH_SIZE]
        messages = build_value_normalization_messages(batch, field_specs_by_name)
        response = provider.generate(messages, schema, settings)
        content = response.content
        if not content:
            raise ValueError("LLM value normalization returned an empty response.")
        parsed = parse_llm_value_normalization_response(content)
        tasks_by_id = {task.item_id: task for task in batch}
        for item in parsed.items:
            task = tasks_by_id.get(item.item_id)
            if task is None:
                continue
            if float(item.confidence or 0.0) < 0.65:
                continue
            if is_blank_llm_normalized_value(item.normalized_value) and is_blank_llm_normalized_value(item.value_code):
                continue
            normalizations[(task.field_name, normalization_value_key(task.raw_value))] = item
        if progress_callback is not None:
            progress_callback(min(offset + len(batch), len(tasks)), len(tasks), "normalizing_values")

    return apply_llm_value_normalizations(
        responses=responses,
        field_specs_by_name=field_specs_by_name,
        normalizations=normalizations,
    )


def collect_llm_value_normalization_tasks(
    *,
    responses: list[ExtractionResponse],
    field_specs_by_name: dict[str, FieldSpec],
) -> list[ExcelValueNormalizationTask]:
    tasks: list[ExcelValueNormalizationTask] = []
    seen: set[tuple[str, str]] = set()
    for response in responses:
        for result in response.results:
            field_spec = field_specs_by_name.get(result.field_name)
            if field_spec is None:
                continue
            if not should_attempt_llm_value_normalization(result, field_spec):
                continue
            key = (result.field_name, normalization_value_key(result.final_value))
            if key in seen:
                continue
            seen.add(key)
            tasks.append(
                ExcelValueNormalizationTask(
                    item_id=f"v{len(tasks) + 1}",
                    field_name=result.field_name,
                    raw_value=result.final_value,
                    review_reasons=list(result.review_reasons or []),
                )
            )
    return tasks


def should_attempt_llm_value_normalization(result: ExtractionFieldResult, field_spec: FieldSpec) -> bool:
    if result.status == "not_found" or not result.needs_review:
        return False
    if is_blank_llm_normalized_value(result.final_value):
        return False
    reasons = set(result.review_reasons or [])
    if reasons & LLM_VALUE_NORMALIZATION_REASONS:
        return True
    return field_spec.field_type in {"radio", "dropdown", "checkbox", "yesno"} and "choice" in " ".join(reasons)


def build_value_normalization_llm_settings(llm_settings: dict[str, Any]) -> dict[str, Any]:
    settings = dict(llm_settings)
    settings["temperature"] = 0
    settings["think"] = False
    try:
        current_max_tokens = int(settings.get("max_tokens", 2048) or 2048)
    except (TypeError, ValueError):
        current_max_tokens = 2048
    settings["max_tokens"] = min(max(current_max_tokens, 768), 2048)
    return settings


def build_value_normalization_messages(
    tasks: list[ExcelValueNormalizationTask],
    field_specs_by_name: dict[str, FieldSpec],
) -> list[dict[str, str]]:
    items: list[dict[str, Any]] = []
    for task in tasks:
        field_spec = field_specs_by_name[task.field_name]
        items.append(
            {
                "item_id": task.item_id,
                "field_name": task.field_name,
                "form_name": clean_metadata_text(field_spec.form_name),
                "field_label": clean_metadata_text(field_spec.field_label),
                "field_type": clean_metadata_text(field_spec.field_type),
                "text_validation": clean_metadata_text(field_spec.text_validation),
                "text_validation_min": clean_metadata_text(field_spec.text_validation_min),
                "text_validation_max": clean_metadata_text(field_spec.text_validation_max),
                "choices": [
                    {"code": choice.code, "label": choice.label}
                    for choice in field_spec.choices_options
                ],
                "raw_value": task.raw_value,
                "review_reasons": task.review_reasons,
            }
        )

    return [
        {
            "role": "system",
            "content": (
                "You normalize invalid imported Excel values for REDCap fields. "
                "Return JSON only. Be conservative: do not invent information. "
                "Only fix spelling, language, abbreviations, category labels/codes, decimal commas, and date formats."
            ),
        },
        {
            "role": "user",
            "content": json.dumps(
                {
                    "task": "Convert each raw_value to a value accepted by the REDCap field definition.",
                    "rules": [
                        "Return one response item per input item_id.",
                        "If the value cannot be converted with high confidence, set normalized_value to null and value_code to null.",
                        "For radio/dropdown/yesno fields, value_code must be one of the provided choice codes.",
                        "For checkbox fields, value_code may be an array of provided choice codes.",
                        "For choice fields, normalized_value should be the exact matching choice label when possible.",
                        "For date fields, normalized_value must be YYYY-MM-DD.",
                        "For integer/number fields, normalized_value must be numeric and respect min/max when provided.",
                        "Do not change a value just to make it fit if the original value is likely from a wrong Excel column.",
                    ],
                    "items": items,
                    "response_schema": {
                        "items": [
                            {
                                "item_id": "same id from input",
                                "normalized_value": "string | number | boolean | array | null",
                                "value_code": "string | array | null",
                                "confidence": 0.0,
                                "reason": "short reason",
                            }
                        ]
                    },
                },
                ensure_ascii=False,
            ),
        },
    ]


def parse_llm_value_normalization_response(content: str) -> LLMValueNormalizationResponse:
    normalized = normalize_json_response_text(content)
    try:
        return LLMValueNormalizationResponse.model_validate_json(normalized)
    except ValidationError as json_exc:
        candidates = [normalized]
        first_object = normalized.find("{")
        last_object = normalized.rfind("}")
        if first_object != -1 and last_object > first_object:
            candidates.append(normalized[first_object : last_object + 1].strip())
        first_array = normalized.find("[")
        last_array = normalized.rfind("]")
        if first_array != -1 and last_array > first_array:
            candidates.append(normalized[first_array : last_array + 1].strip())
        for candidate in candidates:
            try:
                raw = json.loads(candidate)
            except json.JSONDecodeError:
                continue
            if isinstance(raw, dict) and isinstance(raw.get("items"), list):
                return LLMValueNormalizationResponse.model_validate(raw)
            if isinstance(raw, list):
                return LLMValueNormalizationResponse.model_validate({"items": raw})
        raise ValueError(f"invalid_json:{json_exc}") from json_exc


def apply_llm_value_normalizations(
    *,
    responses: list[ExtractionResponse],
    field_specs_by_name: dict[str, FieldSpec],
    normalizations: dict[tuple[str, str], LLMValueNormalizationItem],
) -> int:
    applied = 0
    for response in responses:
        for result in response.results:
            field_spec = field_specs_by_name.get(result.field_name)
            if field_spec is None:
                continue
            key = (result.field_name, normalization_value_key(result.final_value))
            item = normalizations.get(key)
            if item is None:
                continue
            if apply_llm_value_normalization_to_result(result, field_spec, item):
                applied += 1
    return applied


def apply_llm_value_normalization_to_result(
    result: ExtractionFieldResult,
    field_spec: FieldSpec,
    item: LLMValueNormalizationItem,
) -> bool:
    candidate_input = item.normalized_value
    if field_spec.field_type in {"radio", "dropdown", "checkbox", "yesno"} and not is_blank_llm_normalized_value(
        item.value_code
    ):
        candidate_input = item.value_code
    normalized_result = build_field_result_from_excel_value(
        field_spec=field_spec,
        field_name=result.field_name,
        raw_value=candidate_input,
        row_index=0,
        column_name="llm_value_normalization",
        value_transform="direct",
    )
    if normalized_result is None or normalized_result.needs_review:
        return False

    original_value = result.final_value
    result.status = "found"
    result.confidence = min(max(float(item.confidence or 0.0), 0.0), 1.0)
    result.final_value = normalized_result.final_value
    result.final_value_code = normalized_result.final_value_code
    result.needs_review = False
    result.review_reasons = []
    result.selection_reason = "excel_import_llm_value_normalization"
    result.candidates = [
        {
            "value_raw": scalar_to_text(original_value),
            "value_normalized": normalized_result.final_value,
            "value_code": normalized_result.final_value_code,
            "confidence": result.confidence,
            "evidence": f"LLM value normalization: {item.reason or 'converted to REDCap-compatible value'}",
            "context_date": None,
            "context_label": "excel_import_llm_value_normalization",
        },
        *(result.candidates or []),
    ]
    return True


def normalization_value_key(value: Any) -> str:
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    return scalar_to_text(value)


def is_blank_llm_normalized_value(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    if isinstance(value, list):
        return not value
    return False


def build_row_response(
    *,
    project_name: str,
    row: dict[str, Any],
    row_index: int,
    mappings: list[ExcelColumnMapping],
    field_specs_by_name: dict[str, FieldSpec],
) -> ExtractionResponse:
    results: list[ExtractionFieldResult] = []
    for mapping in mappings:
        raw_value = row.get(mapping.column_name)
        if raw_value is None:
            continue
        transformed_value = transform_excel_value(raw_value, mapping.value_transform)
        if transformed_value is None:
            continue
        field_spec = field_specs_by_name.get(mapping.field_name)
        field_result = build_field_result_from_excel_value(
            field_spec=field_spec,
            field_name=mapping.field_name,
            raw_value=transformed_value,
            row_index=row_index,
            column_name=mapping.column_name,
            value_transform=mapping.value_transform,
        )
        if field_result is not None:
            results.append(field_result)
    return ExtractionResponse(project_name=project_name, results=results)


def resolve_patient_identity(
    *,
    rows: list[dict[str, Any]],
    patient_key: str,
    patient_id_column: str,
    requested_patient_mode: str | None,
    requested_identifier_type: str | None,
    column_mappings: list[ExcelColumnMapping],
) -> PatientIdentity:
    headers = list(rows[0].keys()) if rows else []
    record_id_column = resolve_optional_column(headers, RECORD_ID_COLUMN_CANDIDATES)
    tc_column = resolve_optional_column(headers, TC_COLUMN_CANDIDATES) or find_mapped_column(
        column_mappings,
        "tc_kimlik_no",
    )

    record_id = first_non_empty_value(rows, record_id_column) if record_id_column else None
    tc_identity_no = normalize_tc_identity_no(first_non_empty_value(rows, tc_column)) if tc_column else None

    if requested_identifier_type == "record_id" and not record_id:
        record_id = patient_key
    if requested_identifier_type == "tc_kimlik_no" and not tc_identity_no:
        tc_identity_no = normalize_tc_identity_no(patient_key)
    if not tc_identity_no and looks_like_tc_identity_no(patient_key):
        tc_identity_no = normalize_tc_identity_no(patient_key)

    if requested_patient_mode == "new":
        return PatientIdentity(
            patient_mode="new",
            identifier_type=None,
            identifier_value=None,
            tc_identity_no=tc_identity_no,
        )
    if record_id:
        return PatientIdentity(
            patient_mode="existing",
            identifier_type="record_id",
            identifier_value=record_id,
            tc_identity_no=tc_identity_no,
        )
    if tc_identity_no:
        return PatientIdentity(
            patient_mode="auto",
            identifier_type="tc_kimlik_no",
            identifier_value=tc_identity_no,
            tc_identity_no=tc_identity_no,
        )
    return PatientIdentity(
        patient_mode="auto",
        identifier_type=None,
        identifier_value=patient_key,
        tc_identity_no=None,
    )


def build_row_group_key(
    *,
    row: dict[str, Any],
    primary_column: str | None,
    headers: list[str],
    column_mappings: list[ExcelColumnMapping],
) -> PatientGroupKey | None:
    if primary_column:
        primary_value = normalize_identifier_value(row.get(primary_column))
        if primary_value:
            return PatientGroupKey(
                key=f"{normalize_match_key(primary_column)}:{primary_value}",
                value=primary_value,
                source_column=primary_column,
            )

    record_id_column = resolve_optional_column(headers, RECORD_ID_COLUMN_CANDIDATES)
    record_id_value = normalize_identifier_value(row.get(record_id_column)) if record_id_column else None
    if record_id_column and record_id_value:
        return PatientGroupKey(
            key=f"record_id:{record_id_value}",
            value=record_id_value,
            source_column=record_id_column,
        )

    tc_column = resolve_optional_column(headers, TC_COLUMN_CANDIDATES) or find_mapped_column(
        column_mappings,
        "tc_kimlik_no",
    )
    tc_value = normalize_tc_identity_no(row.get(tc_column)) if tc_column else None
    if tc_column and tc_value:
        return PatientGroupKey(
            key=f"tc_kimlik_no:{tc_value}",
            value=tc_value,
            source_column=tc_column,
        )

    fallback_column = resolve_optional_column(headers, IDENTIFIER_COLUMN_CANDIDATES)
    fallback_value = normalize_identifier_value(row.get(fallback_column)) if fallback_column else None
    if fallback_column and fallback_value:
        return PatientGroupKey(
            key=f"{normalize_match_key(fallback_column)}:{fallback_value}",
            value=fallback_value,
            source_column=fallback_column,
        )

    return None


def ensure_tc_identifier_result(
    *,
    response: ExtractionResponse,
    tc_identity_no: str,
    field_specs_by_name: dict[str, FieldSpec],
) -> None:
    if any(result.field_name == "tc_kimlik_no" for result in response.results):
        return
    field_spec = field_specs_by_name.get("tc_kimlik_no")
    response.results.insert(
        0,
        ExtractionFieldResult(
            field_name="tc_kimlik_no",
            form_name=field_spec.form_name if field_spec else None,
            status="found",
            cardinality="single",
            selection_rule="excel_import",
            candidates=[
                {
                    "value_raw": tc_identity_no,
                    "value_normalized": tc_identity_no,
                    "value_code": None,
                    "confidence": 1.0,
                    "evidence": "Excel patient identifier column",
                    "context_date": None,
                    "context_label": "excel_import",
                }
            ],
            confidence=1.0,
            final_value=tc_identity_no,
            selection_reason="excel_import",
            needs_review=False,
            review_reasons=[],
        ),
    )


def build_field_result_from_excel_value(
    *,
    field_spec: FieldSpec | None,
    field_name: str,
    raw_value: Any,
    row_index: int,
    column_name: str,
    value_transform: str = "direct",
) -> ExtractionFieldResult | None:
    value = normalize_cell_value(raw_value)
    if value is None:
        return None

    form_name = field_spec.form_name if field_spec else None
    field_type = field_spec.field_type if field_spec else "text"
    final_value = value
    final_value_code: str | None = None
    status = "found"
    needs_review = False
    review_reasons: list[str] = []

    if field_spec is not None:
        value_for_field = value
        if str(field_spec.text_validation or "").strip().lower().startswith("date"):
            normalized_date = normalize_excel_date_value(value)
            if normalized_date is not None:
                value_for_field = normalized_date
        final_value, final_value_code, choice_review_reasons = normalize_choice_value(value_for_field, field_spec)
        if choice_review_reasons:
            status = "uncertain"
            needs_review = True
            review_reasons.extend(choice_review_reasons)
        validation_reasons = validate_excel_value(final_value, field_spec)
        if validation_reasons:
            status = "uncertain"
            needs_review = True
            review_reasons.extend(validation_reasons)

    candidate = {
        "value_raw": scalar_to_text(value),
        "value_normalized": final_value,
        "value_code": final_value_code,
        "confidence": 1.0,
        "evidence": build_excel_evidence(row_index=row_index, column_name=column_name, value_transform=value_transform),
        "context_date": None,
        "context_label": "excel_import",
    }
    result = ExtractionFieldResult(
        field_name=field_name,
        form_name=form_name,
        status=status,
        cardinality="multiple" if field_type == "checkbox" else "single",
        selection_rule="excel_import",
        candidates=[candidate],
        confidence=1.0,
        final_value=final_value,
        final_value_code=final_value_code,
        selection_reason="excel_import",
        needs_review=needs_review,
        review_reasons=dedupe_strings(review_reasons),
    )
    if field_spec is not None:
        original_final_value = result.final_value
        result = post_process_result(result, field_spec)
        if result.final_value != original_final_value and result.candidates:
            result.candidates[0]["value_normalized"] = result.final_value
    return result


def transform_excel_value(value: Any, value_transform: str | None) -> Any:
    normalized_transform = normalize_value_transform(value_transform)
    normalized_value = normalize_cell_value(value)
    if normalized_value is None or normalized_transform == "direct":
        return normalized_value
    text = scalar_to_text(normalized_value).strip()
    if not text:
        return None
    parts = [part for part in re.split(r"\s+", text) if part]
    if normalized_transform == "first_2":
        return text[:2]
    if normalized_transform == "name_family":
        return parts[-1] if len(parts) >= 2 else None
    if normalized_transform == "name_given":
        if len(parts) >= 2:
            return " ".join(parts[:-1])
        return text
    if normalized_transform == "name_family_first_2":
        return parts[-1][:2] if len(parts) >= 2 else None
    if normalized_transform == "name_given_first_2":
        given_name = " ".join(parts[:-1]) if len(parts) >= 2 else text
        return given_name[:2]
    return normalized_value


def build_excel_evidence(*, row_index: int, column_name: str, value_transform: str | None) -> str:
    transform = normalize_value_transform(value_transform)
    evidence = f"Excel row {row_index}, column {column_name}"
    if transform != "direct":
        evidence = f"{evidence}, transform {transform}"
    return evidence


def normalize_choice_value(value: Any, field_spec: FieldSpec) -> tuple[Any, str | None, list[str]]:
    if field_spec.field_type == "yesno":
        yesno_code = normalize_yesno_code(value)
        if yesno_code is None:
            return value, None, ["choice_code_unmapped"]
        return ("Evet" if yesno_code == "1" else "Hayir"), yesno_code, []

    if field_spec.field_type not in {"radio", "dropdown", "checkbox"}:
        return value, None, []

    if not field_spec.choices_options:
        return value, None, []

    if field_spec.field_type == "checkbox":
        parts = split_checkbox_value(value)
        if not parts:
            return value, None, ["choice_code_unmapped"]
        labels: list[str] = []
        codes: list[str] = []
        unknown_parts: list[str] = []
        for part in parts:
            label, code = match_choice(part, field_spec)
            if code is None:
                unknown_parts.append(str(part))
                continue
            labels.append(label if label is not None else str(part))
            codes.append(code)
        if unknown_parts:
            return parts, json.dumps(codes, ensure_ascii=False) if codes else None, ["choice_code_unmapped"]
        return labels, json.dumps(codes, ensure_ascii=False), []

    label, code = match_choice(value, field_spec)
    if code is None:
        return value, None, ["choice_code_unmapped"]
    return label if label is not None else value, code, []


def match_choice(value: Any, field_spec: FieldSpec) -> tuple[str | None, str | None]:
    value_text = scalar_to_text(value)
    normalized_value = normalize_match_key(value_text)
    for option in field_spec.choices_options:
        if value_text == option.code or normalized_value == normalize_match_key(option.code):
            return option.label, option.code
    for option in field_spec.choices_options:
        if normalized_value == normalize_match_key(option.label):
            return option.label, option.code
    return None, None


def validate_excel_value(value: Any, field_spec: FieldSpec) -> list[str]:
    validation = str(field_spec.text_validation or "").strip().lower()
    reasons: list[str] = []
    if not validation:
        return reasons

    if validation in {"integer"} and not is_integer_like(value):
        reasons.append("integer_validation_failed")
    if validation in {"number", "float"} and not is_number_like(value):
        reasons.append("number_validation_failed")
    if validation.startswith("date") and not is_date_like(value):
        reasons.append("date_validation_failed")

    if validation in {"integer", "number", "float"} and is_number_like(value):
        numeric_value = to_float(value)
        min_value = parse_optional_float(field_spec.text_validation_min)
        max_value = parse_optional_float(field_spec.text_validation_max)
        if min_value is not None and numeric_value is not None and numeric_value < min_value:
            reasons.append("min_validation_failed")
        if max_value is not None and numeric_value is not None and numeric_value > max_value:
            reasons.append("max_validation_failed")

    return reasons


def resolve_patient_id_column(headers: list[str], requested_column: str | None = None) -> str:
    if requested_column:
        for header in headers:
            if header == requested_column:
                return header
        requested_key = normalize_match_key(requested_column)
        for header in headers:
            if normalize_match_key(header) == requested_key:
                return header
        raise ValueError(f"Patient identifier column not found: {requested_column}")

    detected = resolve_optional_column(headers, IDENTIFIER_COLUMN_CANDIDATES)
    if detected:
        return detected
    raise ValueError("Excel sheet must include a patient identifier column such as record_id or tc_kimlik_no.")


def resolve_optional_column(headers: list[str], candidate_names: list[str]) -> str | None:
    header_by_key = {normalize_match_key(header): header for header in headers}
    for candidate in candidate_names:
        header = header_by_key.get(normalize_match_key(candidate))
        if header:
            return header
    return None


def resolve_identifier_type(patient_id_column: str, requested_type: str | None = None) -> str:
    if requested_type in {"record_id", "tc_kimlik_no"}:
        return requested_type
    normalized_column = normalize_match_key(patient_id_column)
    if normalized_column in {
        normalize_match_key("tc_kimlik_no"),
        normalize_match_key("tc_identity_no"),
        normalize_match_key("tc_identity"),
        normalize_match_key("tc_no"),
    }:
        return "tc_kimlik_no"
    return "record_id"


def normalize_patient_mode(patient_mode: str | None) -> str:
    return "new" if patient_mode == "new" else "existing"


def build_queue_label(
    *,
    rows: list[dict[str, Any]],
    label_column: str | None,
    patient_id_column: str,
    patient_key: str,
) -> str:
    if label_column:
        for row in rows:
            label = normalize_identifier_value(row.get(label_column))
            if label:
                return label
    return f"{patient_id_column}: {patient_key}"


def sample_column_values(rows: list[dict[str, Any]], header: str, limit: int = 6) -> list[Any]:
    values: list[Any] = []
    seen: set[str] = set()
    for row in rows:
        value = normalize_cell_value(row.get(header))
        if value is None:
            continue
        key = scalar_to_text(value)
        if key in seen:
            continue
        seen.add(key)
        values.append(value)
        if len(values) >= limit:
            break
    return values


def find_mapped_column(mappings: list[ExcelColumnMapping], field_name: str) -> str | None:
    for mapping in mappings:
        if mapping.field_name == field_name:
            return mapping.column_name
    return None


def first_non_empty_value(rows: list[dict[str, Any]], column_name: str | None) -> str | None:
    if not column_name:
        return None
    for row in rows:
        value = normalize_identifier_value(row.get(column_name))
        if value:
            return value
    return None


def normalize_tc_identity_no(value: Any) -> str | None:
    if value in {None, ""}:
        return None
    digits = "".join(ch for ch in scalar_to_text(value) if ch.isdigit())
    if len(digits) != 11:
        return None
    return digits


def looks_like_tc_identity_no(value: Any) -> bool:
    return normalize_tc_identity_no(value) is not None


def first_non_empty_row_index(rows: list[tuple[Any, ...]]) -> int | None:
    for index, row in enumerate(rows):
        if any(normalize_cell_value(cell) is not None for cell in row):
            return index
    return None


def normalize_headers(raw_headers: tuple[Any, ...]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for value in raw_headers:
        header = scalar_to_text(normalize_cell_value(value)).strip()
        if not header:
            headers.append("")
            continue
        count = seen.get(header, 0)
        seen[header] = count + 1
        headers.append(header if count == 0 else f"{header}_{count + 1}")
    return headers


def normalize_cell_value(value: Any) -> Any:
    if is_missing_scalar(value):
        return None
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        stripped = value.strip()
        if stripped.lower() in NULL_TEXT_VALUES:
            return None
        return stripped
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def normalize_identifier_value(value: Any) -> str | None:
    value = normalize_cell_value(value)
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = scalar_to_text(value).strip()
    return text or None


def scalar_to_text(value: Any) -> str:
    if is_missing_scalar(value):
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def clean_metadata_text(value: Any) -> str:
    if is_missing_scalar(value):
        return ""
    return scalar_to_text(value).strip()


def is_missing_scalar(value: Any) -> bool:
    return value is None or (isinstance(value, float) and math.isnan(value))


def normalize_search_text(value: Any) -> str:
    text = scalar_to_text(value).strip().lower()
    text = text.replace("ı", "i")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def tokenize_match_text(value: Any) -> set[str]:
    tokens = {
        token
        for token in normalize_search_text(value).split()
        if len(token) > 1 and token not in GENERIC_FIELD_TOKENS
    }
    expanded = set(tokens)
    for token in tokens:
        expanded.update(SYNONYM_LOOKUP.get(token, set()))
    return expanded


def acronym_for_text(value: Any) -> str:
    return "".join(token[0] for token in normalize_search_text(value).split() if token)


def normalize_match_key(value: Any) -> str:
    text = scalar_to_text(value).strip().lower()
    text = text.replace("ı", "i")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", text)


def unique_lookup(pairs: Any) -> dict[str, str]:
    lookup: dict[str, str] = {}
    ambiguous: set[str] = set()
    for key, value in pairs:
        if not key:
            continue
        if key in lookup and lookup[key] != value:
            ambiguous.add(key)
            continue
        lookup[key] = value
    for key in ambiguous:
        lookup.pop(key, None)
    return lookup


def find_embedded_field_name(header: str, normalized_field_names: dict[str, str]) -> str | None:
    chunks = re.split(r"[|:/()\[\]{}]", header)
    for chunk in chunks:
        field_name = normalized_field_names.get(normalize_match_key(chunk))
        if field_name:
            return field_name
    return None


def normalize_yesno_code(value: Any) -> str | None:
    text = normalize_match_key(value)
    if text in {"1", "true", "yes", "y", "evet", "e", "var"}:
        return "1"
    if text in {"0", "false", "no", "n", "hayir", "h", "yok"}:
        return "0"
    return None


def split_checkbox_value(value: Any) -> list[Any]:
    if isinstance(value, list):
        return [item for item in value if normalize_cell_value(item) is not None]
    text = scalar_to_text(value).strip()
    if not text:
        return []
    separators = r"[;|\n]"
    if not re.search(separators, text) and "," in text:
        separators = r","
    return [part.strip() for part in re.split(separators, text) if part.strip()]


def is_integer_like(value: Any) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, int):
        return True
    if isinstance(value, float):
        return value.is_integer()
    text = scalar_to_text(value).strip()
    return bool(re.fullmatch(r"[-+]?\d+", text))


def is_number_like(value: Any) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    return to_float(value) is not None


def to_float(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(scalar_to_text(value).strip().replace(",", "."))
    except ValueError:
        return None


def is_date_like(value: Any) -> bool:
    return normalize_excel_date_value(value) is not None


def normalize_excel_date_value(value: Any) -> str | None:
    if isinstance(value, (datetime, date)):
        date_value = value.date() if isinstance(value, datetime) else value
        return date_value.isoformat()
    text = scalar_to_text(value).strip()
    if not text:
        return None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return text
    for pattern in (
        "%d.%m.%Y",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%d.%m.%y",
        "%d/%m/%y",
        "%d-%m-%y",
        "%m/%d/%Y",
        "%m-%d-%Y",
        "%m/%d/%y",
        "%m-%d-%y",
    ):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    return None


def parse_optional_float(value: Any) -> float | None:
    if value in {None, ""}:
        return None
    try:
        return float(str(value).strip().replace(",", "."))
    except ValueError:
        return None


def dedupe_strings(values: list[str]) -> list[str]:
    seen: set[str] = set()
    deduped: list[str] = []
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        deduped.append(value)
    return deduped
